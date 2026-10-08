from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file, flash
from google import genai
from dotenv import load_dotenv
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from functools import wraps
import mysql.connector
from mysql.connector import pooling
from questions import (
    ROLES,
    ROLE_DESCRIPTIONS,
    generate_questions_batch,
    next_question_batch,
    generate_resume_questions_ai,
)
import os
import traceback
import json
import re
import time
import uuid
import hashlib
from io import BytesIO
import PyPDF2
import docx
from report_pdf import generate_interview_pdf

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "dev-secret-key-change-in-production")

# Always reload templates from disk on each request so that a running server
# instantly serves updated templates (e.g. the resume interview upload UI)
# instead of a stale, cached, blank page.
app.jinja_env.auto_reload = True
app.config["TEMPLATES_AUTO_RELOAD"] = True


@app.after_request
def no_cache_html(response):
    # Prevent browsers from showing a stale cached version of HTML pages.
    if "text/html" in response.headers.get("Content-Type", ""):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

# ==========================================
# DATABASE CONNECTION POOL
# ==========================================

db_config = {
    "host": os.getenv("DB_HOST", "localhost"),
    "user": os.getenv("DB_USER", "root"),
    "password": os.getenv("DB_PASSWORD", ""),
    "database": os.getenv("DB_NAME", "ai_voice_coach"),
}

db_pool = None

def get_db():
    global db_pool
    if db_pool is None:
        try:
            db_pool = pooling.MySQLConnectionPool(
                pool_name="aicoach_pool",
                pool_size=5,
                **db_config
            )
        except mysql.connector.Error as e:
            print("Database pool error:", e)
            return None
    try:
        return db_pool.get_connection()
    except mysql.connector.Error as e:
        print("Database connection error:", e)
        return None


def _ensure_interview_columns(cursor):
    """Add any missing report columns to the `interviews` table without wiping
    existing data. MySQL's `ALTER TABLE` does not support `ADD COLUMN IF NOT
    EXISTS`, so we check information_schema first. Reused each startup so an
    existing deployment migrates cleanly on restart."""
    try:
        cursor.execute("""
            SELECT COLUMN_NAME FROM information_schema.COLUMNS
            WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = 'interviews'
        """)
        existing = {row[0] for row in cursor.fetchall()}
    except mysql.connector.Error as e:
        print("Column inspection error:", e)
        return

    additions = [
        ("communication", "INT NULL"),
        ("confidence", "INT NULL"),
        ("clarity", "INT NULL"),
        ("grammar", "INT NULL"),
        ("structure", "INT NULL"),
        ("relevance", "INT NULL"),
        ("details", "JSON NULL"),
        ("strengths", "JSON NULL"),
        ("improvements", "JSON NULL"),
        ("role", "VARCHAR(100) NULL"),
        ("experience", "VARCHAR(50) NULL"),
        ("difficulty", "VARCHAR(20) NULL"),
        ("overall_feedback", "TEXT NULL"),
        ("session_type", "VARCHAR(20) NULL"),
    ]
    for col, definition in additions:
        if col not in existing:
            try:
                cursor.execute("ALTER TABLE interviews ADD COLUMN %s %s" % (col, definition))
            except mysql.connector.Error as e:
                print("Add column %s error: %s" % (col, e))


def _ensure_user_columns(cursor):
    """Add any missing role / status columns to the `users` table without
    wiping existing data. Existing users automatically default to role='user'
    and is_active=1, preserving backward compatibility."""
    try:
        cursor.execute("""
            SELECT COLUMN_NAME FROM information_schema.COLUMNS
            WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = 'users'
        """)
        existing = {row[0] for row in cursor.fetchall()}
    except mysql.connector.Error as e:
        print("User column inspection error:", e)
        return

    additions = [
        ("role", "VARCHAR(20) NOT NULL DEFAULT 'user'"),
        ("is_active", "TINYINT(1) NOT NULL DEFAULT 1"),
    ]
    for col, definition in additions:
        if col not in existing:
            try:
                cursor.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, definition))
            except mysql.connector.Error as e:
                print("Add user column %s error: %s" % (col, e))


def init_db():
    conn = get_db()
    if not conn:
        print("WARNING: Could not connect to database. Auth features will not work.")
        return
    try:
        cursor = conn.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INT AUTO_INCREMENT PRIMARY KEY,
                full_name VARCHAR(100) NOT NULL,
                email VARCHAR(150) NOT NULL UNIQUE,
                password_hash VARCHAR(255) NOT NULL,
                role VARCHAR(20) NOT NULL DEFAULT 'user',
                is_active TINYINT(1) NOT NULL DEFAULT 1,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS interviews (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                interview_type VARCHAR(50) NOT NULL DEFAULT 'practice',
                overall_score INT DEFAULT 0,
                questions_count INT DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
            )
        """)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS resume_interview_state (
                user_id INT PRIMARY KEY,
                questions JSON NOT NULL,
                analysis JSON NULL,
                role VARCHAR(100) NULL,
                experience VARCHAR(50) NULL,
                count INT DEFAULT 0,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                                ON UPDATE CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
            )
        """)
        _ensure_interview_columns(cursor)
        _ensure_user_columns(cursor)
        conn.commit()
        print("Database tables initialized.")
    except mysql.connector.Error as e:
        print("Database init error:", e)
    finally:
        cursor.close()
        conn.close()


# ==========================================
# AUTH HELPERS
# ==========================================

def safe_next_url(url):
    if not url:
        return ""
    if url.startswith("/") and not url.startswith("//") and "\\" not in url:
        return url
    return ""


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login", next=request.path))
        # Re-check the user still exists and is active on every request so a
        # deactivated/deleted account is blocked immediately.
        user = get_current_user()
        if not user:
            session.clear()
            return redirect(url_for("login", next=request.path))
        if not user.get("is_active"):
            session.clear()
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


def get_current_user():
    if "user_id" not in session:
        return None
    conn = get_db()
    if not conn:
        return None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT id, full_name, email, role, is_active, created_at FROM users WHERE id = %s", (session["user_id"],))
        user = cursor.fetchone()
        return user
    except mysql.connector.Error:
        return None
    finally:
        cursor.close()
        conn.close()


def user_context():
    user = get_current_user()
    return {"user": user}


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login", next=request.path))
        user = get_current_user()
        if not user:
            session.clear()
            return redirect(url_for("login", next=request.path))
        if not user.get("is_active"):
            session.clear()
            return redirect(url_for("login"))
        if user.get("role") != "admin":
            flash("You do not have permission to access the Admin Dashboard.")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return decorated


# ==========================================
# RESUME INTERVIEW STATE (server-side, per-user)
# Stored in MySQL (NOT the cookie session) so it survives refresh / re-entry /
# new tab across the same account without bloating the ~4KB session cookie or
# making another Gemini question-generation request.
# ==========================================

def save_resume_interview_state(user_id, questions, analysis, role, experience, count):
    conn = get_db()
    if not conn:
        return False
    cursor = None
    try:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO resume_interview_state
                (user_id, questions, analysis, role, experience, count)
            VALUES (%s, %s, %s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                questions = VALUES(questions),
                analysis = VALUES(analysis),
                role = VALUES(role),
                experience = VALUES(experience),
                count = VALUES(count)
        """, (user_id,
              json.dumps(questions, default=str),
              json.dumps(analysis, default=str) if analysis else None,
              role, experience, count))
        conn.commit()
        return True
    except Exception as e:
        print("save_resume_interview_state error:", e)
        return False
    finally:
        if cursor is not None:
            try: cursor.close()
            except Exception: pass
        try: conn.close()
        except Exception: pass


def get_resume_interview_state(user_id):
    conn = get_db()
    if not conn:
        return None
    cursor = None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT questions, analysis FROM resume_interview_state WHERE user_id = %s", (user_id,))
        row = cursor.fetchone()
        if not row:
            return None
        try:
            questions = json.loads(row["questions"])
        except Exception:
            return None
        if not isinstance(questions, list) or len(questions) == 0:
            return None
        analysis = None
        if row.get("analysis"):
            try:
                analysis = json.loads(row["analysis"])
            except Exception:
                analysis = None
        return {"questions": questions, "analysis": analysis}
    except Exception as e:
        print("get_resume_interview_state error:", e)
        return None
    finally:
        if cursor is not None:
            try: cursor.close()
            except Exception: pass
        try: conn.close()
        except Exception: pass


def delete_resume_interview_state(user_id):
    conn = get_db()
    if not conn:
        return False
    cursor = None
    try:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM resume_interview_state WHERE user_id = %s", (user_id,))
        conn.commit()
        return True
    except Exception as e:
        print("delete_resume_interview_state error:", e)
        return False
    finally:
        if cursor is not None:
            try: cursor.close()
            except Exception: pass
        try: conn.close()
        except Exception: pass


def get_user_statistics(user_id):
    conn = get_db()
    if not conn:
        return {
            "total_interviews": 0, "average_score": None, "best_score": None,
            "current_streak": 0, "total_questions": 0,
            "communication": None, "confidence": None, "clarity": None,
            "grammar": None, "structure": None, "relevance": None
        }
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT COUNT(*) AS total, ROUND(AVG(overall_score)) AS avg_score, "
            "MAX(overall_score) AS best_score, COALESCE(SUM(questions_count), 0) AS total_questions "
            "FROM interviews WHERE user_id = %s",
            (user_id,)
        )
        row = cursor.fetchone()
        total = row["total"] or 0
        avg_score = int(row["avg_score"]) if row["avg_score"] is not None else None
        best_score = int(row["best_score"]) if row["best_score"] is not None else None
        total_questions = int(row["total_questions"] or 0)

        cursor.execute(
            """SELECT
                ROUND(AVG(communication)) AS communication,
                ROUND(AVG(confidence)) AS confidence,
                ROUND(AVG(clarity)) AS clarity,
                ROUND(AVG(grammar)) AS grammar,
                ROUND(AVG(structure)) AS structure,
                ROUND(AVG(relevance)) AS relevance
               FROM interviews
               WHERE user_id = %s AND communication IS NOT NULL""",
            (user_id,)
        )
        cats = cursor.fetchone()
        category_avgs = {}
        if cats:
            for k, v in cats.items():
                if v is not None:
                    category_avgs[k] = int(v)

        cursor.execute(
            "SELECT DISTINCT DATE(created_at) AS d FROM interviews WHERE user_id = %s ORDER BY d ASC",
            (user_id,)
        )
        dates = [str(r["d"]) for r in cursor.fetchall()]

        streak = 0
        if dates:
            from datetime import date, timedelta
            today = date.today()
            unique_dates = sorted(set(dates), reverse=True)
            date_set = set(unique_dates)
            check = today
            if str(check) not in date_set:
                check -= timedelta(days=1)
            while str(check) in date_set:
                streak += 1
                check -= timedelta(days=1)

        return {
            "total_interviews": total,
            "average_score": avg_score,
            "best_score": best_score,
            "current_streak": streak,
            "total_questions": total_questions,
            **category_avgs
        }
    except mysql.connector.Error as e:
        print("Statistics query error:", e)
        return {
            "total_interviews": 0, "average_score": None, "best_score": None,
            "current_streak": 0, "total_questions": 0,
            "communication": None, "confidence": None, "clarity": None,
            "grammar": None, "structure": None, "relevance": None
        }
    finally:
        cursor.close()
        conn.close()


def get_user_recent_interviews(user_id, limit=20):
    conn = get_db()
    if not conn:
        return []
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """SELECT id, interview_type, overall_score, questions_count, created_at,
                      role, difficulty,
                      communication, confidence, clarity, grammar, structure, relevance
               FROM interviews WHERE user_id = %s ORDER BY created_at DESC LIMIT %s""",
            (user_id, limit)
        )
        return cursor.fetchall()
    except mysql.connector.Error as e:
        print("Recent interviews query error:", e)
        return []
    finally:
        cursor.close()
        conn.close()


# ==========================================
# ADMIN HELPERS
# ==========================================

ADMIN_CATEGORIES = [
    ("communication", "Communication"),
    ("confidence", "Confidence"),
    ("clarity", "Clarity"),
    ("grammar", "Grammar"),
    ("structure", "Structure"),
    ("relevance", "Relevance"),
]

# Classify each interview into practice/mock/resume. New sessions store the
# `session_type` column; older rows (NULL) are inferred from the distinctive
# overall_feedback text each session type writes.
SESSION_TYPE_EXPR = """
    COALESCE(session_type,
        CASE
            WHEN interview_type IN ('resume', 'resume-based') THEN 'resume'
            WHEN overall_feedback LIKE 'Mock interview%' THEN 'mock'
            WHEN overall_feedback LIKE 'Practice session%' THEN 'practice'
            WHEN interview_type IN ('behavioral', 'communication') THEN 'practice'
            ELSE 'practice'
        END
    )
"""


def _safe_int(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def get_admin_summary_stats():
    """Aggregate statistics across the whole system for the admin dashboard."""
    conn = get_db()
    empty = {
        "total_users": 0,
        "total_practice": 0,
        "total_mock": 0,
        "total_resume": 0,
        "total_interviews": 0,
        "avg_overall": None,
        "highest_overall": None,
        "lowest_overall": None,
        "avg_communication": None,
        "avg_confidence": None,
        "avg_clarity": None,
        "avg_grammar": None,
        "avg_structure": None,
        "avg_relevance": None,
    }
    if not conn:
        return empty
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT COUNT(*) AS c FROM users")
        total_users = _safe_int(cursor.fetchone()["c"])

        cursor.execute(
            "SELECT COUNT(*) AS c FROM interviews WHERE %s = 'practice'" % SESSION_TYPE_EXPR
        )
        total_practice = _safe_int(cursor.fetchone()["c"])

        cursor.execute(
            "SELECT COUNT(*) AS c FROM interviews WHERE %s = 'mock'" % SESSION_TYPE_EXPR
        )
        total_mock = _safe_int(cursor.fetchone()["c"])

        cursor.execute(
            "SELECT COUNT(*) AS c FROM interviews WHERE %s = 'resume'" % SESSION_TYPE_EXPR
        )
        total_resume = _safe_int(cursor.fetchone()["c"])

        cursor.execute("SELECT COUNT(*) AS c FROM interviews")
        total_interviews = _safe_int(cursor.fetchone()["c"])

        row = {"avg_overall": None}
        cursor.execute("SELECT ROUND(AVG(overall_score)) AS avg_score, "
                       "MAX(overall_score) AS max_score, MIN(overall_score) AS min_score "
                       "FROM interviews")
        r = cursor.fetchone()
        if r and r["avg_score"] is not None:
            row["avg_overall"] = _safe_int(r["avg_score"])
        row["highest_overall"] = _safe_int(r["max_score"]) if r and r["max_score"] is not None else None
        row["lowest_overall"] = _safe_int(r["min_score"]) if r and r["min_score"] is not None else None

        cat_avgs = {}
        for key, _label in ADMIN_CATEGORIES:
            cursor.execute(
                "SELECT ROUND(AVG(%s)) AS v FROM interviews WHERE %s IS NOT NULL" % (key, key)
            )
            r = cursor.fetchone()
            cat_avgs[key] = _safe_int(r["v"]) if r and r["v"] is not None else None

        return {
            "total_users": total_users,
            "total_practice": total_practice,
            "total_mock": total_mock,
            "total_resume": total_resume,
            "total_interviews": total_interviews,
            "avg_overall": row["avg_overall"],
            "highest_overall": row["highest_overall"],
            "lowest_overall": row["lowest_overall"],
            **cat_avgs,
        }
    except mysql.connector.Error as e:
        print("Admin summary stats error:", e)
        return empty
    finally:
        cursor.close()
        conn.close()


def get_admin_all_users():
    """Return all users with per-user interview aggregates for user management."""
    conn = get_db()
    if not conn:
        return []
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """SELECT u.id, u.full_name, u.email, u.role, u.is_active, u.created_at,
                      COUNT(i.id) AS total_interviews,
                      SUM(CASE WHEN i.id IS NOT NULL AND """ + SESSION_TYPE_EXPR + """ = 'practice' THEN 1 ELSE 0 END) AS practice_count,
                      SUM(CASE WHEN i.id IS NOT NULL AND """ + SESSION_TYPE_EXPR + """ = 'mock' THEN 1 ELSE 0 END) AS mock_count,
                      ROUND(AVG(i.overall_score)) AS avg_score
               FROM users u
               LEFT JOIN interviews i ON i.user_id = u.id
               GROUP BY u.id, u.full_name, u.email, u.role, u.is_active, u.created_at
               ORDER BY u.created_at DESC"""
        )
        return cursor.fetchall()
    except mysql.connector.Error as e:
        print("Admin all users error:", e)
        return []
    finally:
        cursor.close()
        conn.close()


def get_user_by_id(user_id):
    conn = get_db()
    if not conn:
        return None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT id, full_name, email, role, is_active, created_at FROM users WHERE id = %s",
            (user_id,)
        )
        return cursor.fetchone()
    except mysql.connector.Error as e:
        print("Get user by id error:", e)
        return None
    finally:
        cursor.close()
        conn.close()


def get_admin_user_detail(user_id):
    """Full performance detail for a single user as seen by an admin."""
    conn = get_db()
    if not conn:
        return {"stats": None, "interviews": []}
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """SELECT id, full_name, email, role, is_active, created_at FROM users WHERE id = %s""",
            (user_id,)
        )
        user = cursor.fetchone()
        if not user:
            return {"stats": None, "interviews": []}

        cursor.execute(
            """SELECT COUNT(*) AS total,
                      SUM(CASE WHEN """ + SESSION_TYPE_EXPR + """ = 'practice' THEN 1 ELSE 0 END) AS practice_count,
                      SUM(CASE WHEN """ + SESSION_TYPE_EXPR + """ = 'mock' THEN 1 ELSE 0 END) AS mock_count,
                      ROUND(AVG(overall_score)) AS avg_score,
                      MAX(overall_score) AS best_score
               FROM interviews WHERE user_id = %s""",
            (user_id,)
        )
        agg = cursor.fetchone() or {}

        cursor.execute(
            """SELECT id, interview_type, overall_score, created_at,
                      communication, confidence, clarity, grammar, structure, relevance
               FROM interviews WHERE user_id = %s
               ORDER BY created_at DESC LIMIT 20""",
            (user_id,)
        )
        interviews = cursor.fetchall()

        stats = {
            "total_interviews": _safe_int((agg or {}).get("total")),
            "practice_count": _safe_int((agg or {}).get("practice_count")),
            "mock_count": _safe_int((agg or {}).get("mock_count")),
            "avg_score": _safe_int((agg or {}).get("avg_score")) if (agg or {}).get("avg_score") is not None else None,
            "best_score": _safe_int((agg or {}).get("best_score")) if (agg or {}).get("best_score") is not None else None,
            "latest_score": (interviews[0]["overall_score"] if interviews else None),
        }

        strongest = None
        weakest = None
        if interviews:
            latest = interviews[0]
            scores = {key: latest.get(key) for key, _label in ADMIN_CATEGORIES}
            present = [(key, val) for key, val in scores.items() if val is not None]
            if present:
                strongest = max(present, key=lambda kv: kv[1])
                weakest = min(present, key=lambda kv: kv[1])

        return {
            "user": user,
            "stats": stats,
            "interviews": interviews,
            "strongest": strongest,
            "weakest": weakest,
        }
    except mysql.connector.Error as e:
        print("Admin user detail error:", e)
        return {"stats": None, "interviews": []}
    finally:
        cursor.close()
        conn.close()


def get_admin_analytics():
    """System-wide analytics including per-category averages, common weak/strong
    areas, interview activity over time and new-user registration trend."""
    conn = get_db()
    if not conn:
        return None
    com_weak = (None, 0)
    com_strong = (None, 0)
    try:
        cursor = conn.cursor(dictionary=True)

        cat_avgs = {}
        for key, _label in ADMIN_CATEGORIES:
            cursor.execute(
                "SELECT ROUND(AVG(%s)) AS v, COUNT(*) AS n FROM interviews WHERE %s IS NOT NULL" % (key, key)
            )
            r = cursor.fetchone()
            cat_avgs[key] = {
                "avg": _safe_int(r["v"]) if r and r["v"] is not None else None,
                "count": _safe_int(r["n"]),
            }

        # Most common weakest / strongest area across all per-question records
        # is approximated by aggregating each interview's own weakest/strongest
        # category (so every interview contributes exactly one vote each).
        cursor.execute(
            """SELECT id, communication, confidence, clarity, grammar, structure, relevance
               FROM interviews WHERE communication IS NOT NULL OR confidence IS NOT NULL
               OR clarity IS NOT NULL OR grammar IS NOT NULL OR structure IS NOT NULL
               OR relevance IS NOT NULL"""
        )
        rows = cursor.fetchall()
        weak_counts = {key: 0 for key, _label in ADMIN_CATEGORIES}
        strong_counts = {key: 0 for key, _label in ADMIN_CATEGORIES}
        for row in rows:
            scores = {key: row.get(key) for key, _label in ADMIN_CATEGORIES}
            present = [(key, val) for key, val in scores.items() if val is not None]
            if present:
                max_kv = max(present, key=lambda kv: kv[1])
                min_kv = min(present, key=lambda kv: kv[1])
                strong_counts[max_kv[0]] += 1
                weak_counts[min_kv[0]] += 1
        com_weak = (None, 0)
        com_strong = (None, 0)
        if weak_counts:
            com_weak = max(weak_counts.items(), key=lambda kv: kv[1])
        if strong_counts:
            com_strong = max(strong_counts.items(), key=lambda kv: kv[1])

        # Interview activity over the last 30 days.
        cursor.execute(
            """SELECT DATE(created_at) AS d, COUNT(*) AS n
               FROM interviews
               WHERE created_at >= DATE_SUB(CURDATE(), INTERVAL 30 DAY)
               GROUP BY DATE(created_at) ORDER BY d ASC"""
        )
        activity_rows = cursor.fetchall()

        # New user registrations over the last 30 days.
        cursor.execute(
            """SELECT DATE(created_at) AS d, COUNT(*) AS n
               FROM users
               WHERE created_at >= DATE_SUB(CURDATE(), INTERVAL 30 DAY)
               GROUP BY DATE(created_at) ORDER BY d ASC"""
        )
        reg_rows = cursor.fetchall()

        return {
            "cat_avgs": cat_avgs,
            "common_weak": com_weak,
            "common_strong": com_strong,
            "weak_dist": weak_counts,
            "strong_dist": strong_counts,
            "activity": activity_rows,
            "registrations": reg_rows,
        }
    except mysql.connector.Error as e:
        print("Admin analytics error:", e)
        return None
    finally:
        cursor.close()
        conn.close()


def get_recent_registrations(limit=5):
    conn = get_db()
    if not conn:
        return []
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT id, full_name, email, created_at FROM users ORDER BY created_at DESC LIMIT %s",
            (limit,)
        )
        return cursor.fetchall()
    except mysql.connector.Error as e:
        print("Recent registrations error:", e)
        return []
    finally:
        cursor.close()
        conn.close()


def get_recent_interviews_all(limit=5):
    conn = get_db()
    if not conn:
        return []
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """SELECT i.id, i.interview_type, i.overall_score, i.created_at, u.full_name, u.email
               FROM interviews i JOIN users u ON u.id = i.user_id
               ORDER BY i.created_at DESC LIMIT %s""",
            (limit,)
        )
        return cursor.fetchall()
    except mysql.connector.Error as e:
        print("Recent interviews all error:", e)
        return []
    finally:
        cursor.close()
        conn.close()


# ==========================================
# GEMINI API
# ==========================================

api_key = os.getenv("GEMINI_API_KEY")
if not api_key:
    print("WARNING: GEMINI_API_KEY not found in .env")

# Configure a reasonable request timeout so a slow/hung Gemini response fails
# cleanly instead of depending on httpx's short default. HttpOptions.timeout is
# expressed in milliseconds (120s here, enough for long resume analysis).
client = (
    genai.Client(
        api_key=api_key,
        http_options=genai.types.HttpOptions(timeout=120000),
    )
    if api_key
    else None
)

# Short-timeout client for interactive answer evaluation (mock + resume
# interviews). A single attempt with a 30s socket timeout keeps the
# "Evaluating your answer..." state bounded: if Gemini is slow/hung, the
# /analyze route fails fast and degrades to safe local scoring instead of
# leaving the UI spinning for minutes on retries.
evaluate_client = (
    genai.Client(
        api_key=api_key,
        http_options=genai.types.HttpOptions(timeout=30000),
    )
    if api_key
    else None
)

# Maximum retries for transient Gemini errors (rate limits / server hiccups).
GEMINI_MAX_RETRIES = 3
# Network/socket errors (e.g. WinError 10054) are usually transient — the next
# fresh connection succeeds. We allow at most ONE controlled retry, never more.
NETWORK_MAX_RETRIES = 1


class AIError(Exception):
    """Raised when the AI provider call fails.

    `category` is one of: "quota", "auth", "config", "network", "server".
    """
    def __init__(self, message, category="server"):
        super().__init__(message)
        self.category = category


def _parse_retry_after(error_text):
    """Best-effort parse of Gemini's 'Please retry in Ns' hint -> seconds."""
    m = re.search(r"retry\s+in\s+([0-9.]+)\s*s", str(error_text), re.IGNORECASE)
    if m:
        try:
            return float(m.group(1))
        except ValueError:
            return None
    return None


def _is_network_error(e, final_status, msg):
    """True when the failure is a transport/socket-level error (e.g. WinError
    10054 connection reset), not an HTTP status error. These are usually
    transient and a fresh connection often succeeds. Called only when no
    HTTP status (`.code`) is available from the SDK exception."""
    if final_status is not None:
        return False
    if isinstance(e, (ConnectionError, TimeoutError, OSError)):
        return True
    low = str(msg).lower()
    markers = (
        "10054", "existing connection was forcibly closed", "connection reset",
        "connection aborted", "broken pipe", "ecosystem", "econnreset",
        "econnaborted", "epipe", "winerror", "remote closed", "remote protocol error",
        "remote end closed", "connection refused", "connection closed", "unexpected eof",
        "timed out", "timeout", "read timed out", "the server closed the connection",
    )
    return any(m in low for m in markers)


def _log_ai_call(feature, ok, status=None, attempts=1, detail=""):
    """Safe development logging. Never logs the API key or prompt content."""
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    print(
        f"AI REQUEST: feature={feature} time={ts} "
        f"success={'yes' if ok else 'no'} status={status} attempts={attempts}"
        f"{(' detail=' + detail[:120]) if detail else ''}"
    )


def gemini_generate_text(prompt, model="gemini-3.5-flash", feature="ai", timeout_ms=None):
    """Call Gemini with retry/backoff for transient errors.

    Respects the provider's Retry-After hint. Distinguishes error categories so
    callers can surface the right message. Never fakes a result.
    Logs each request (feature/timestamp/status) without exposing the key.

    Interactive, latency-sensitive calls (answer evaluation) pass a short
    `timeout_ms`: they run on the single-attempt 30s evaluate_client and fail
    fast, so a slow/hung Gemini never leaves the user staring at a loading
    spinner for minutes.
    """
    if not client:
        raise AIError("AI service is not configured. Check GEMINI_API_KEY in .env.", category="config")

    if timeout_ms is not None:
        use_client = evaluate_client or client
        max_attempts = 1
    else:
        use_client = client
        max_attempts = GEMINI_MAX_RETRIES

    last_exc = None
    final_status = None
    attempts = 0
    for attempt in range(1, max_attempts + 1):
        attempts = attempt
        try:
            response = use_client.models.generate_content(model=model, contents=prompt)
            raw = response.text.strip()
            if raw.startswith("```"):
                raw = raw.split("\n", 1)[1]
                if raw.endswith("```"):
                    raw = raw[:-3]
                raw = raw.strip()
            _log_ai_call(feature, True, status=200, attempts=attempts)
            return raw
        except Exception as e:
            last_exc = e
            final_status = getattr(e, "code", None)
            msg = str(e)
            retry_after = _parse_retry_after(msg)
            is_quota = (final_status == 429) or ("RESOURCE_EXHAUSTED" in msg) or ("quota" in msg.lower())
            is_network = _is_network_error(e, final_status, msg)
            # Interactive evaluation calls fail fast with a clear category so
            # the caller can degrade to local scoring instead of waiting.
            if timeout_ms is not None:
                if is_quota:
                    _log_ai_call(feature, False, status=429, attempts=attempts, detail="quota/rate limit")
                    hint = f" Please retry in about {retry_after:.0f}s." if retry_after else ""
                    raise AIError(
                        "Gemini API rate limit reached. Please check your API usage or try again later." + hint,
                        category="quota",
                    ) from e
                if final_status in (400, 401, 403) or "API key" in msg or "apikey" in msg.lower():
                    _log_ai_call(feature, False, status=final_status, attempts=attempts, detail="auth")
                    raise AIError("AI API key is invalid or unavailable.", category="auth") from e
                _log_ai_call(feature, False, status=final_status, attempts=attempts, detail="evaluation timeout/failure")
                raise AIError(
                    "AI analysis timed out or failed: " + msg[:200],
                    category="network" if is_network else "server",
                ) from e
            # Respect provider's Retry-After, falling back to a small backoff.
            if retry_after:
                wait = min(retry_after + 1, 30)
            else:
                wait = 2 ** attempt

            # Quota (429) is a hard account limit: retrying repeatedly just
            # re-hits the exhausted quota and extends the wait. Surface it once.
            if is_quota:
                _log_ai_call(feature, False, status=429, attempts=attempts, detail="quota/rate limit")
                hint = f" Please retry in about {retry_after:.0f}s." if retry_after else ""
                raise AIError(
                    "Gemini API rate limit reached. Please check your API usage or try again later." + hint,
                    category="quota",
                ) from e
            # HTTP 5xx provider errors: retry a few times with small backoff.
            if final_status in (500, 502, 503, 504):
                if attempt < GEMINI_MAX_RETRIES:
                    print(f"Gemini transient error (status={final_status}) on attempt {attempt}; "
                          f"retrying in {wait:.0f}s: {msg[:150]}")
                    time.sleep(wait)
                    continue
                _log_ai_call(feature, False, status=final_status, attempts=attempts, detail="provider 5xx")
                raise AIError(f"AI service is temporarily unavailable (HTTP {final_status}). Please try again soon.",
                              category="server") from e

            # Network/socket errors (e.g. WinError 10054) are transient — a fresh
            # connection usually succeeds. Allow at most ONE controlled retry.
            if is_network:
                if attempt <= NETWORK_MAX_RETRIES:
                    print(f"Gemini network error on attempt {attempt}; retrying once in {wait:.0f}s: {msg[:150]}")
                    time.sleep(wait)
                    continue
                _log_ai_call(feature, False, status=None, attempts=attempts, detail="network/connection")
                raise AIError(
                    "AI service connection was interrupted. Please try again in a moment.", category="network"
                ) from e
            if final_status in (400, 401, 403) or "API key" in msg or "apikey" in msg.lower():
                _log_ai_call(feature, False, status=final_status, attempts=attempts, detail="auth")
                raise AIError("AI API key is invalid or unavailable.", category="auth") from e
            if final_status == 404 or "404" in msg or "not found" in msg.lower():
                _log_ai_call(feature, False, status=final_status, attempts=attempts, detail="model")
                raise AIError(f"AI model '{model}' is not available for this API account.", category="config") from e
            _log_ai_call(feature, False, status=final_status, attempts=attempts, detail="server")
            raise AIError(f"AI generation failed: {msg[:300]}", category="server") from e

    _log_ai_call(feature, False, status=final_status, attempts=attempts, detail="exhausted")
    raise AIError(f"AI generation failed after retries: {str(last_exc)}", category="server")


# ==========================================
# AUTH ROUTES
# ==========================================

@app.route("/signup", methods=["GET", "POST"])
def signup():
    if "user_id" in session:
        return redirect(url_for("dashboard"))

    next_url = safe_next_url(request.args.get("next"))

    if request.method == "GET":
        return render_template("signup.html", next=next_url)

    full_name = request.form.get("full_name", "").strip()
    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")
    confirm = request.form.get("confirm_password", "")

    errors = []
    if not full_name:
        errors.append("Please enter your full name.")
    elif len(full_name) < 2 or len(full_name) > 100:
        errors.append("Name must be between 2 and 100 characters.")

    if not email:
        errors.append("Please enter your email address.")
    elif not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email):
        errors.append("Please enter a valid email address.")

    if not password:
        errors.append("Please enter a password.")
    elif len(password) < 8:
        errors.append("Password must contain at least 8 characters.")

    if password != confirm:
        errors.append("Passwords do not match.")

    if errors:
        return render_template("signup.html", errors=errors, full_name=full_name, email=email, next=next_url)

    conn = get_db()
    if not conn:
        return render_template("signup.html", errors=["Unable to connect to the database. Please try again later."], full_name=full_name, email=email, next=next_url)

    try:
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM users WHERE email = %s", (email,))
        if cursor.fetchone():
            return render_template("signup.html", errors=["An account with this email already exists."], full_name=full_name, email=email, next=next_url)

        password_hash = generate_password_hash(password)
        cursor.execute(
            "INSERT INTO users (full_name, email, password_hash) VALUES (%s, %s, %s)",
            (full_name, email, password_hash)
        )
        conn.commit()
        return render_template("login.html", success="Account created successfully. Please sign in.", auth_required=bool(next_url), next=next_url)
    except mysql.connector.Error as e:
        print("Signup error:", e)
        return render_template("signup.html", errors=["Unable to create account. Please try again later."], full_name=full_name, email=email, next=next_url)
    finally:
        cursor.close()
        conn.close()


@app.route("/login", methods=["GET", "POST"])
def login():
    if "user_id" in session:
        return redirect(url_for("dashboard"))

    next_url = safe_next_url(request.args.get("next") or request.form.get("next") or "")
    auth_required = bool(next_url)

    if request.method == "GET":
        return render_template("login.html", auth_required=auth_required, next=next_url)

    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")

    if not email or not password:
        return render_template("login.html", errors=["Please enter both email and password."], email=email, auth_required=auth_required, next=next_url)

    conn = get_db()
    if not conn:
        return render_template("login.html", errors=["Unable to connect to the database. Please try again later."], email=email, auth_required=auth_required, next=next_url)

    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT id, full_name, email, password_hash, role, is_active FROM users WHERE email = %s", (email,))
        user = cursor.fetchone()

        if not user or not check_password_hash(user["password_hash"], password):
            return render_template("login.html", errors=["Invalid email or password."], email=email, auth_required=auth_required, next=next_url)

        if not user.get("is_active", 1):
            return render_template("login.html", errors=["Your account has been deactivated. Please contact an administrator."], email=email, auth_required=auth_required, next=next_url)

        session["user_id"] = user["id"]
        session["user_name"] = user["full_name"]
        session["user_email"] = user["email"]
        session["user_role"] = user["role"]
        session["role"] = user["role"]
        if next_url:
            return redirect(next_url)
        return redirect(url_for("dashboard"))
    except mysql.connector.Error as e:
        print("Login error:", e)
        return render_template("login.html", errors=["Unable to sign in. Please try again later."], email=email, auth_required=auth_required, next=next_url)
    finally:
        cursor.close()
        conn.close()


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("home"))


# ==========================================
# PAGE ROUTES
# ==========================================

@app.route("/")
def home():
    stats = {"total_interviews": 0, "average_score": None, "best_score": None, "current_streak": 0}
    recent_interviews = []
    ctx = user_context()
    if ctx.get("user"):
        stats = get_user_statistics(ctx["user"]["id"])
        recent_interviews = get_user_recent_interviews(ctx["user"]["id"], limit=5)
    return render_template("index.html", stats=stats, recent_interviews=recent_interviews, **ctx)


@app.route("/practice")
def practice():
    interview_type = request.args.get("type", "hr")
    return render_template("practice.html", interview_type=interview_type, **user_context())


@app.route("/mock-interview")
@login_required
def mock_interview():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    # The mock interview persists its completed session to the /save-interview
    # endpoint under the logged-in user's id; require login so the save can
    # never silently fail with 401 and the interview is always recorded for
    # the SAME user the Dashboard / Reports / Improvement Plan read.
    return render_template("interview.html", user_id=user["id"], **user_context())


@app.route("/dashboard")
@login_required
def dashboard():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    stats = get_user_statistics(user["id"])
    recent_interviews = get_user_recent_interviews(user["id"], limit=20)
    return render_template("dashboard.html", stats=stats, recent_interviews=recent_interviews, **user_context())


@app.route("/reports")
@login_required
def reports():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    interviews = get_user_recent_interviews(user["id"], limit=500)
    stats = get_user_statistics(user["id"])
    return render_template("report.html", interviews=interviews, stats=stats, **user_context())


@app.route("/reports/<int:interview_id>")
@login_required
def report_detail(interview_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    conn = get_db()
    if not conn:
        return render_template("report_detail.html", interview=None, error="Database unavailable.", **user_context())
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """SELECT id, interview_type, overall_score, questions_count, created_at,
                      communication, confidence, clarity, grammar, structure, relevance,
                      details, strengths, improvements, role, experience, difficulty, overall_feedback
               FROM interviews WHERE id = %s AND user_id = %s""",
            (interview_id, user["id"])
        )
        interview = cursor.fetchone()
        if not interview:
            return render_template("report_detail.html", interview=None, error="Report not found.", **user_context())

        # Decode JSON columns into safe Python structures for the template.
        def _load_json(value):
            if not value:
                return []
            try:
                parsed = json.loads(value)
                return parsed if isinstance(parsed, list) else []
            except (ValueError, TypeError):
                return []

        interview["details"] = _load_json(interview.get("details"))
        interview["strengths"] = _load_json(interview.get("strengths"))
        interview["improvements"] = _load_json(interview.get("improvements"))
        # Normalize detail items so the template can rely on known keys.
        for item in interview["details"]:
            if isinstance(item, dict):
                item.setdefault("question", "")
                item.setdefault("answer", "")
                item.setdefault("score", None)
                item.setdefault("category", "General")
                item.setdefault("feedback", "")
                item.setdefault("better_answer", "")
                item.setdefault("scores", {})
        return render_template("report_detail.html", interview=interview, **user_context())
    except mysql.connector.Error as e:
        print("Report detail error:", e)
        return render_template("report_detail.html", interview=None, error="Could not load report.", **user_context())
    finally:
        cursor.close()
        conn.close()


def _load_json_list(value):
    """Decode a JSON column into a list (used by the PDF download route).
    Mirrors the decoding the report detail view performs so both consumers
    rely on the same data shape."""
    if not value:
        return []
    try:
        parsed = json.loads(value)
        return parsed if isinstance(parsed, list) else []
    except (ValueError, TypeError):
        return []


_REPORT_COLUMNS = """id, interview_type, overall_score, questions_count, created_at,
                      communication, confidence, clarity, grammar, structure, relevance,
                      details, strengths, improvements, role, experience, difficulty, overall_feedback"""


@app.route("/reports/<int:interview_id>/pdf")
@login_required
def report_download_pdf(interview_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    conn = get_db()
    if not conn:
        return "Database is unavailable. Please try again.", 500
    cursor = None
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT %s FROM interviews WHERE id = %%s AND user_id = %%s" % _REPORT_COLUMNS,
            (interview_id, user["id"])
        )
        interview = cursor.fetchone()
    except mysql.connector.Error as e:
        print("Report PDF download error:", e)
        return "Could not load the report. Please try again.", 500
    finally:
        if cursor is not None:
            cursor.close()
        conn.close()

    # The query already scopes to the logged-in user, so this covers both
    # "report does not exist" and "report belongs to another user".
    if not interview:
        return "Report not found.", 404

    interview["details"] = _load_json_list(interview.get("details"))
    interview["strengths"] = _load_json_list(interview.get("strengths"))
    interview["improvements"] = _load_json_list(interview.get("improvements"))
    for item in interview["details"]:
        if isinstance(item, dict):
            item.setdefault("question", "")
            item.setdefault("answer", "")
            item.setdefault("score", None)
            item.setdefault("category", "General")
            item.setdefault("feedback", "")
            item.setdefault("better_answer", "")

    try:
        pdf_buffer = generate_interview_pdf(interview, user)
    except Exception as e:
        print("PDF generation error:", e)
        return "Could not generate the PDF report. Please try again.", 500

    return send_file(
        pdf_buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name="interview_report_%d.pdf" % interview_id,
    )


# ==========================================
# INTERVIEW QUESTIONS (EXPANDED BANK + AI)
# ==========================================

@app.route("/api/questions")
def get_all_questions():
    interview_type = request.args.get("type", "hr")
    difficulty = request.args.get("difficulty", "beginner")
    try:
        count = int(request.args.get("count", 10))
    except (TypeError, ValueError):
        count = 10
    count = max(1, min(count, 200))
    role = request.args.get("role") or None
    experience = request.args.get("experience") or None
    result = generate_questions_batch(interview_type, difficulty, count, role=role, experience=experience)
    return jsonify({
        "questions": result["questions"],
        "type": interview_type,
        "difficulty": difficulty,
        "count": len(result["questions"]),
        "source": result["source"],
    })


@app.route("/api/question")
def get_question():
    interview_type = request.args.get("type", "hr")
    difficulty = request.args.get("difficulty", "beginner")
    role = request.args.get("role") or None
    result = generate_questions_batch(interview_type, difficulty, 1, role=role)
    question = result["questions"][0] if result["questions"] else "Tell me about yourself."
    return jsonify({"question": question, "type": interview_type, "source": result["source"]})


@app.route("/api/generate-questions", methods=["POST"])
def generate_questions():
    data = request.get_json(silent=True) or {}
    interview_type = data.get("interview_type", "general")
    role = data.get("role", "software_developer")
    experience = data.get("experience", "fresher")
    difficulty = data.get("difficulty", "beginner")
    try:
        count = int(data.get("count", 10))
    except (TypeError, ValueError):
        count = 10
    count = max(1, min(count, 200))
    history = data.get("history", [])

    previous_questions = []
    for item in history:
        if isinstance(item, dict):
            q = item.get("question")
            if q:
                previous_questions.append(q)
        elif isinstance(item, str) and item.strip():
            previous_questions.append(item)

    try:
        result = generate_questions_batch(
            interview_type, difficulty, count,
            role=role, experience=experience, previous_questions=previous_questions,
        )
        questions = result["questions"]
        if not questions:
            return jsonify({"success": False, "error": "Unable to generate interview questions. Please try again later.", "error_type": "server"}), 500
        payload = {
            "success": True,
            "questions": questions,
            "count": len(questions),
            "source": result["source"],
        }
        if result["warning"]:
            payload["warning"] = result["warning"]
        return jsonify(payload)
    except Exception as e:
        print("AI question generation error:", type(e).__name__, str(e)[:300])
        return jsonify({"success": False, "error": "AI question generation failed. Please try again later.", "error_type": "server"}), 500


@app.route("/api/next-question", methods=["POST"])
def next_adaptive_question():
    data = request.get_json(silent=True) or {}
    interview_type = data.get("interview_type", "general")
    role = data.get("role", "software_developer")
    experience = data.get("experience", "fresher")
    difficulty = data.get("difficulty", "beginner")
    history = data.get("history", [])
    question_number = data.get("question_number", 1)
    total_questions = data.get("total_questions", 10)

    try:
        result = next_question_batch(
            interview_type, difficulty, role=role, experience=experience, history=history,
        )
        payload = {"success": True, "question": result["question"]}
        if result["warning"]:
            payload["warning"] = result["warning"]
        return jsonify(payload)
    except Exception as e:
        print("Adaptive question error:", type(e).__name__, str(e)[:300])
        return jsonify({"success": False, "error": "Could not generate next question.", "error_type": "server"}), 500


@app.route("/api/roles")
def get_roles():
    return jsonify({"roles": ROLES})


# ==========================================
# AI ANSWER ANALYSIS
# ==========================================

FILLER_WORDS_LIST = [
    "um", "umm", "uh", "like", "actually",
    "basically", "you know", "so", "I mean",
    "kind of", "sort of"
]

CATEGORY_KEYS = ["communication", "confidence", "clarity", "grammar", "structure", "relevance"]


SCORE_WEIGHTS = {
    "communication": 0.20,
    "confidence": 0.15,
    "clarity": 0.15,
    "grammar": 0.10,
    "structure": 0.20,
    "relevance": 0.20,
}


NON_ANSWER_PHRASES = {
    "i dont know",
    "i do not know",
    "dont know",
    "do not know",
    "i dont know the answer",
    "i do not know the answer",
    "i dont know the answer to this question",
    "i do not know the answer to this question",
    "sorry i dont know",
    "sorry i do not know",
    "sorry i dont know the answer",
    "sorry i do not know the answer",
    "no idea",
    "i have no idea",
    "i have absolutely no idea",
    "i cannot answer",
    "i cant answer",
    "cannot answer",
    "i cannot answer the question",
    "i cant answer the question",
    "no answer",
    "i have no answer",
    "i dont have an answer",
    "i do not have an answer",
    "i dont have the answer",
    "i do not have the answer",
    "not sure",
    "i am not sure",
    "im not sure",
    "i am unsure",
    "skip",
    "skip this question",
}


_NON_ANSWER_FILLERS = {
    "sorry", "i", "im", "am", "so", "please", "just",
    "um", "uh", "umm", "uhm", "uhh", "hmm", "well", "like",
}


def normalize_answer(text):
    text = (text or "").lower().replace("\u2019", "'")
    contractions = {
        "don't": "dont", "can't": "cant", "cannot": "cannot",
        "i'm": "im", "i've": "ive", "i'll": "ill", "i'd": "id",
        "isn't": "isnt", "aren't": "arent", "won't": "wont",
        "wouldn't": "wouldnt", "couldn't": "couldnt", "doesn't": "doesnt",
        "didn't": "didnt", "that's": "thats", "what's": "whats",
        "it's": "its", "let's": "lets",
    }
    for old, new in contractions.items():
        text = text.replace(old, new)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def is_non_answer(text):
    """True when the answer is empty/whitespace, only filler words, or an
    explicit refusal (e.g. "I don't know"). Short but meaningful answers
    (e.g. "Python is an interpreted programming language.") are NOT treated
    as non-answers and still go through AI evaluation."""
    normalized = normalize_answer(text)
    if not normalized:
        return True
    if normalized in NON_ANSWER_PHRASES:
        return True
    for phrase in sorted(NON_ANSWER_PHRASES, key=len, reverse=True):
        if phrase in normalized:
            remainder = [
                w for w in normalized.replace(phrase, " ").split()
                if w not in _NON_ANSWER_FILLERS
            ]
            if not remainder:
                return True
    return False


def non_answer_result(message):
    zero_scores = {key: 0 for key in CATEGORY_KEYS}
    return {
        "success": True,
        "overall_score": 0,
        "scores": zero_scores,
        "clarity": 0,
        "communication": 0,
        "confidence": 0,
        "grammar": 0,
        "relevance": 0,
        "answer_structure": 0,
        "strengths": [],
        "areas_to_improve": ["No meaningful answer was provided."],
        "improvements": ["No meaningful answer was provided."],
        "suggestions": [
            "Try to answer the question with your understanding.",
            "Even if you are unsure, explain your approach or thoughts."
        ],
        "feedback": message,
        "ai_feedback": message,
        "better_answer": "",
        "interview_tip": "",
        "ai_error": "",
        "feedback_note": message,
        "filler_data": {"total": 0, "rate": "0.0", "words": {}, "word_count": 0},
        "star_analysis": {"s": False, "t": False, "a": False, "r": False},
        "star_score": 0,
        "ai_available": False,
        "non_answer": True,
    }


def calculate_overall_score(scores):
    score = sum(scores.get(k, 0) * w for k, w in SCORE_WEIGHTS.items())
    return round(score)


def validate_scores(scores):
    validated = {}
    for key in CATEGORY_KEYS:
        val = scores.get(key)
        # Robust to numeric strings ("75") the AI sometimes returns.
        try:
            num = float(val)
        except (TypeError, ValueError):
            num = None
        if num is not None and 0 <= num <= 100:
            validated[key] = int(round(num))
        else:
            # No artificial minimum/default. Missing or invalid categories are 0
            # so they never inflate the overall score.
            validated[key] = 0
    return validated


def calculate_filler_words(text):
    lower_text = text.lower()
    filler_counts = {}
    total = 0
    for filler in FILLER_WORDS_LIST:
        count = len(re.findall(r'\b' + re.escape(filler) + r'\b', lower_text))
        if count > 0:
            filler_counts[filler] = count
            total += count
    word_count = len(text.split())
    rate = round((total / word_count) * 100, 1) if word_count > 0 else 0
    return {"total": total, "rate": rate, "words": filler_counts, "word_count": word_count}


def _as_text(value):
    return str(value).strip() if value else ""


def _as_string_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(x).strip() for x in value if str(x) and str(x).strip()]
    if isinstance(value, str):
        stripped = value.strip()
        return [stripped] if stripped else []
    return []


def _parse_ai_json(raw, label="AI response"):
    """Robustly parse AI text into a JSON value.

    Tolerates markdown code fences (```json ... ``` or ``` ... ```), BOM,
    smart quotes, and prose wrapped around the object. Raises AIError with the
    real reason and a raw snippet when nothing parses so failures are loud
    instead of being silently replaced.
    """
    if raw is None:
        raise AIError(f"{label} was empty.", category="server")
    text = str(raw).replace("\ufeff", "").strip()
    if not text:
        raise AIError(f"{label} was empty.", category="server")

    attempts = [text]

    # Normalize smart/curly quotes to ASCII so the JSON still parses.
    swapped = (text.replace("\u201c", '"').replace("\u201d", '"')
                   .replace("\u2018", "'").replace("\u2019", "'"))
    if swapped != text:
        attempts.append(swapped)

    # Extract the body of a markdown code block, e.g. ```json\n{...}\n```.
    for block in re.findall(r"```[a-zA-Z]*[^\n]*\n(.*?)```", text, re.S):
        block = block.strip()
        attempts.append(block)
        s, e = block.find("{"), block.rfind("}")
        if s != -1 and e > s:
            attempts.append(block[s:e + 1])

    # First '{' .. last '}' span regardless of surrounding prose.
    s, e = text.find("{"), text.rfind("}")
    if s != -1 and e > s:
        attempts.append(text[s:e + 1])

    for candidate in attempts:
        candidate = candidate.strip()
        if not candidate:
            continue
        try:
            parsed = json.loads(candidate)
        except (ValueError, TypeError):
            continue
        if isinstance(parsed, (dict, list)):
            return parsed

    raise AIError(
        f"{label} was not valid JSON. Raw start: {text[:200]!r}",
        category="server",
    )


def normalize_ai_evaluation(data):
    """Normalize whatever the AI returned into the canonical evaluation shape.

    Accepts either the nested ``scores`` dict or flat top-level category keys
    (``clarity``, ``communication``, ``confidence``, ``grammar``,
    ``relevance``, ``answer_structure``), maps ``weaknesses`` /
    ``improvements`` onto the canonical ``areas_to_improve`` list, and keeps
    the detailed analysis / corrected answer / improved sample answer fields.
    """
    scores_in = data.get("scores", {})
    if not isinstance(scores_in, dict):
        scores_in = {}

    flat_map = {
        "communication": data.get("communication"),
        "confidence": data.get("confidence"),
        "clarity": data.get("clarity"),
        "grammar": data.get("grammar"),
        "relevance": data.get("relevance"),
        "structure": data.get("answer_structure") or data.get("structure"),
    }
    for key, val in flat_map.items():
        if key not in scores_in and val is not None:
            scores_in[key] = val

    scores = validate_scores(scores_in)

    strengths = _as_string_list(data.get("strengths"))
    areas_to_improve = _as_string_list(
        data.get("areas_to_improve")
        or data.get("improvements")
        or data.get("weaknesses")
    )
    suggestions = _as_string_list(
        data.get("suggestions") or data.get("recommendations") or data.get("next_steps")
    )
    feedback = _as_text(data.get("ai_feedback") or data.get("feedback"))

    detailed_in = data.get("detailed_analysis")
    if not isinstance(detailed_in, dict):
        detailed_in = {}
    detailed_analysis = {
        "summary": _as_text(detailed_in.get("summary")),
        "content_analysis": _as_text(detailed_in.get("content_analysis")),
        "technical_accuracy": _as_text(detailed_in.get("technical_accuracy")),
        "communication_analysis": _as_text(detailed_in.get("communication_analysis")),
    }

    improved_sample_answer = _as_text(
        data.get("improved_sample_answer") or data.get("better_answer")
    )
    corrected_answer = _as_text(
        data.get("corrected_answer")
        or data.get("corrected version of the answer")
    )

    speaking_in = data.get("speaking_analysis")
    speaking_analysis = None
    if isinstance(speaking_in, dict):
        cleaned = {
            "filler_words": speaking_in.get("filler_words"),
            "duration": _as_text(speaking_in.get("duration")),
            "speech_rate": speaking_in.get("speech_rate"),
            "pace_feedback": _as_text(speaking_in.get("pace_feedback")),
            "confidence_feedback": _as_text(speaking_in.get("confidence_feedback")),
        }
        if any(v for v in cleaned.values()):
            speaking_analysis = cleaned

    overall = data.get("overall_score")
    if not isinstance(overall, (int, float)) or not (0 <= overall <= 100):
        overall = calculate_overall_score(scores)

    return {
        "overall_score": round(overall),
        "scores": scores,
        "detailed_analysis": detailed_analysis,
        "strengths": strengths,
        "areas_to_improve": areas_to_improve,
        "suggestions": suggestions,
        "corrected_answer": corrected_answer,
        "improved_sample_answer": improved_sample_answer,
        "feedback": feedback,
        "better_answer": improved_sample_answer,
        "interview_tip": _as_text(data.get("interview_tip")),
        "speaking_analysis": speaking_analysis,
    }


def analyze_with_gemini(text, interview_type="general", question="", interview_role=""):
    if not client:
        return {"error": "AI service unavailable. Please check your API configuration."}

    question_context = ""
    if question:
        question_context = f"\nInterview Question: \"{question}\"\n"
    role_context = ""
    if interview_role:
        role_context = f"\nTarget Role: {interview_role}\n"

    prompt = f"""You are an expert AI interview coach. Analyze the following interview answer and return a JSON response.

Interview Type: {interview_type}
{question_context}
{role_context}
Answer: "{text}"

Return ONLY a valid JSON object with EXACTLY this structure:
{{
    "overall_score": <number 0-100>,
    "scores": {{
        "communication": <number 0-100>,
        "confidence": <number 0-100>,
        "clarity": <number 0-100>,
        "grammar": <number 0-100>,
        "relevance": <number 0-100>,
        "structure": <number 0-100>
    }},
    "detailed_analysis": {{
        "summary": "<1-3 sentence overall summary of the answer>",
        "content_analysis": "<what the candidate covered and what important points were missing>",
        "technical_accuracy": "<accuracy and technical depth, judged against the target role>",
        "communication_analysis": "<clarity, structure, and flow of the answer>"
    }},
    "strengths": ["<strength1>", "<strength2>", "<strength3>"],
    "areas_to_improve": ["<improvement1>", "<improvement2>", "<improvement3>"],
    "corrected_answer": "<grammar-corrected, clearer version of the CANDIDATE'S OWN answer. Preserve their main intention and the points they made. Fix grammar and flow only. Do not add skills/concepts they did not mention.>",
    "improved_sample_answer": "<a stronger, more detailed, interview-quality model answer for this question and target role>",
    "interview_tip": "<one useful interview tip>",
    "suggestions": ["<one actionable tip>", "<another actionable tip>"],
    "speaking_analysis": {{
        "filler_words": <0>,
        "duration": "<estimated spoken duration>",
        "speech_rate": <estimated words per minute>,
        "pace_feedback": "<advice on delivery pace>",
        "confidence_feedback": "<advice on sounding more confident>"
    }},
    "ai_feedback": "<2-3 paragraph detailed evaluation>"
}}

Always return these exact field names. "detailed_analysis" and
"speaking_analysis" are JSON objects. "strengths", "areas_to_improve" and
"suggestions" MUST be JSON arrays of strings (never a single string or an
object). "corrected_answer" must be a faithful correction of the candidate's
own words; "improved_sample_answer" is a separate, richer model answer.
Base every field (scores, analysis, corrected and improved answers) on the
ACTUAL interview question, the ACTUAL answer, and the target role.

Category definitions:
- Communication: clarity of expression, vocabulary, flow
- Confidence: assertiveness, tone, directness
- Clarity: how easy it is to understand
- Grammar: grammatical correctness
- Relevance: how well the answer addresses the interview question (if provided)
- Structure: organization, logical flow, completeness

Scoring guidelines (use the FULL 0-100 range; there is NO minimum score):
- 90-100: Excellent and detailed answer
- 75-89: Good answer
- 60-74: Reasonable answer with some weaknesses
- 40-59: Weak or partially relevant answer
- 20-39: Very weak answer
- 1-19: Extremely poor or almost no meaningful answer
- 0: No meaningful answer

Critical rules:
- If the answer refuses to answer or says the candidate does not know
  (e.g. "I don't know", "Sorry, I don't know the answer", "No idea",
  "I cannot answer", "No answer"), the overall_score and EVERY category
  score MUST be exactly 0. Strengths should be empty.
- Do NOT artificially increase scores. If an answer deserves 0, return 0.
- Score in the full 0-100 range. A very weak answer may score as low as
  1-19, never padded to 50.
- Provide all six category scores, each an integer between 0 and 100.

Return ONLY the JSON object, no other text."""

    try:
        raw = gemini_generate_text(prompt, feature="answer_evaluation", timeout_ms=30000)
        try:
            data = _parse_ai_json(raw, label="AI evaluation response")
        except AIError as e:
            # Loud failure: log the real reason and the raw text so the issue
            # is never silently hidden behind a generic "unavailable" message.
            print("EVALUATE PARSE ERROR:", e)
            print("EVALUATE RAW TEXT:", repr(raw[:500]))
            return {"error": str(e)}
        if not isinstance(data, dict):
            print("EVALUATE TYPE ERROR: expected JSON object but got",
                  type(data).__name__, "| raw:", repr(raw[:500]))
            return {"error": "AI evaluation response was not a JSON object."}
        return normalize_ai_evaluation(data)
    except AIError as e:
        # Propagate quota/auth/config errors so the API can surface them clearly.
        raise e
    except Exception as e:
        print("Gemini Error:", type(e).__name__, str(e)[:300])
        raise AIError(f"AI analysis failed: {str(e)[:200]}", category="server") from e


def _fallback_ideal_answer(question, interview_role=""):
    """Static guidance used ONLY when the AI/API genuinely fails to generate
    the ideal answer (quota, config, network). The improved sample answer is
    still produced so the Corrected/Sample sections are never empty."""
    corrected = (
        "No answer was provided. Focus on understanding the key concepts required "
        "to answer this question."
    )
    improved = (
        "Review this topic and practice explaining the main concept, your approach, "
        "and a relevant example."
    )
    tip = (
        "Even when you are unsure, explain what you know and describe how you "
        "would approach the problem."
    )
    return corrected, improved, tip


def generate_ideal_answer(question, interview_role=""):
    """Generate the ideal/correct answer for an interview question via AI.

    Called when the candidate provided no meaningful answer: scores stay 0,
    but the user still receives a model answer to learn from. The candidate's
    text is deliberately NOT sent in this prompt (it is not scored here).
    Raises AIError so the caller can fall back to static guidance when the
    AI provider is unavailable.
    """
    if not client:
        raise AIError(
            "AI service is not configured. Check GEMINI_API_KEY in .env.",
            category="config",
        )

    role_context = ""
    if (interview_role or "").strip():
        role_context = f"\nTarget role: {interview_role.strip()}\n"

    prompt = f"""You are an expert interview coach.

The candidate did not provide a meaningful answer to this interview question.

Interview Question:
{question}
{role_context}
Generate:
1. A short, encouraging explanation of what the candidate should address in their answer ("corrected_answer").
2. A high-quality, correct, professional interview sample answer ("improved_sample_answer") that fully answers the question.

Return ONLY a valid JSON object with EXACTLY this structure:
{{
  "corrected_answer": "<2-3 sentence explanation of what a good response should cover>",
  "improved_sample_answer": "<complete, accurate, polished model answer for the question and target role>",
  "interview_tip": "<one concise tip on how to attempt an answer even when unsure>"
}}

Rules:
- improved_sample_answer must directly and completely answer the interview question.
- It must be clear, accurate, professional, and suitable for the specified target role.
- Do not mention that this answer was AI-generated.
- Do not include meta text such as "Here is a sample answer".

Return ONLY the JSON object, no other text."""

    try:
        raw = gemini_generate_text(prompt, feature="ideal_answer", timeout_ms=30000)
        data = _parse_ai_json(raw, label="Ideal answer response")
    except AIError as e:
        print("IDEAL ANSWER PARSE ERROR:", e)
        print("IDEAL ANSWER RAW TEXT:", repr(raw[:500]) if "raw" in locals() else "")
        raise e
    if not isinstance(data, dict):
        raise AIError("Ideal answer response was not a JSON object.", category="server")

    corrected = _as_text(data.get("corrected_answer"))
    improved = _as_text(data.get("improved_sample_answer"))
    if not corrected and not improved:
        raise AIError(
            "Ideal answer response contained no corrected or improved answer.",
            category="server",
        )
    return {
        "corrected_answer": corrected,
        "improved_sample_answer": improved,
        "interview_tip": _as_text(data.get("interview_tip")),
    }


def detect_star(text):
    lower = text.lower()
    return {
        "s": bool(re.search(r'\b(situation|context|background|when|once|during)\b', lower)),
        "t": bool(re.search(r'\b(task|responsible|needed|had to|goal|objective)\b', lower)),
        "a": bool(re.search(r'\b(action|did|implemented|developed|created|used|approach|decided)\b', lower)),
        "r": bool(re.search(r'\b(result|outcome|achieved|improved|saved|increased|reduced|learned)\b', lower))
    }


def _local_fallback_scores(text, question, star_score, filler_rate):
    """Independent per-category heuristic scores used ONLY when the AI is down.

    Each category is derived from a different signal (question keyword overlap,
    STAR structure, sentence length, filler rate, weak-word usage) so the six
    values are naturally different — never a single shared number. These are
    clearly local approximations, never presented as AI judgment.
    """
    words = re.findall(r"[a-z0-9']+", text.lower())
    word_count = len(words)

    qwords = {w for w in re.findall(r"[a-z0-9']+", (question or "").lower()) if len(w) > 2}
    awords = {w for w in words if len(w) > 2}
    if qwords:
        overlap = len(qwords & awords) / max(len(qwords), 1)
        relevance = round(max(0.0, min(100.0, overlap * 100 + star_score * 0.4)))
    else:
        relevance = round(max(0.0, min(100.0, star_score * 0.5 + word_count * 2.0)))

    structure = round(max(0.0, min(100.0, float(star_score))))

    sentences = [s.strip() for s in re.split(r"[.!?]+", text) if s.strip()]
    avg_words = (word_count / max(len(sentences), 1)) if sentences else word_count
    grammar = round(max(0.0, min(100.0, 80.0 - max(0.0, avg_words - 18.0) * 1.5)))

    clarity = round(max(0.0, min(100.0, 55.0 + (min(word_count, 40) / 40.0) * 35.0 - filler_rate * 0.6)))

    weak_pat = r"\b(maybe|i think|not sure|i guess|kind of|probably|sort of|i don't know)\b"
    weak_count = len(re.findall(weak_pat, text.lower()))
    confidence = round(max(0.0, min(100.0, 58.0 + (min(word_count, 40) / 40.0) * 25.0 - weak_count * 5.0 - filler_rate * 0.3)))

    communication = round(clarity * 0.6 + structure * 0.4)

    return {
        "communication": communication,
        "confidence": confidence,
        "clarity": clarity,
        "grammar": grammar,
        "relevance": relevance,
        "structure": structure,
    }


@app.route("/analyze", methods=["POST"])
def analyze():
    data = request.get_json(silent=True) or {}
    text = data.get("text", "").strip()
    interview_type = data.get("interview_type", "general")
    question = data.get("question", "").strip()
    question_index = data.get("question_index")
    user_id = data.get("user_id")
    interview_role = (data.get("interview_role") or data.get("role") or "").strip()

    print(f"EVALUATE REQUEST: interview_type={interview_type} "
          f"question={question[:80]!r} text_len={len(text)} "
          f"role={interview_role!r} question_index={question_index} user_id={user_id}")

    no_answer = (not text) or is_non_answer(text)

    if no_answer:
        message = (
            "No meaningful answer was provided. Try to attempt the question and "
            "explain your thoughts, even if you are not completely sure."
            if not text
            else "No meaningful answer was provided for this question. Try to "
                 "attempt the question and explain your thoughts, even if you are "
                 "not completely sure."
        )
        result = non_answer_result(message)

        # Even when the user declines to answer, still generate the ideal answer
        # for the CURRENT question so the Corrected Answer / Improved Sample
        # Answer sections are never empty. Scores stay exactly 0. The candidate's
        # non-answer text is NOT sent for scoring.
        ideal_error = ""
        ideal = None
        if question:
            try:
                ideal = generate_ideal_answer(question, interview_role)
            except AIError as e:
                ideal_error = str(e)
                print(f"IDEAL ANSWER AI UNAVAILABLE: {ideal_error}")
            except Exception as e:
                ideal_error = str(e)
                print("IDEAL ANSWER ERROR:", e)
                print(traceback.format_exc())

        if ideal is None:
            corrected_answer, improved_sample_answer, interview_tip = (
                _fallback_ideal_answer(question, interview_role)
            )
        else:
            corrected_answer = ideal.get("corrected_answer", "")
            improved_sample_answer = ideal.get("improved_sample_answer", "")
            interview_tip = (
                ideal.get("interview_tip", "")
                or "Even when you are unsure, explain what you know and describe "
                   "how you would approach the problem."
            )

        result.update({
            "detailed_analysis": "No meaningful answer was provided for this "
                                 "question.",
            "corrected_answer": corrected_answer,
            "improved_sample_answer": improved_sample_answer,
            "better_answer": improved_sample_answer,
            "interview_tip": interview_tip,
            "ai_error": ideal_error,
            "ai_available": ideal is not None,
            "feedback_note": (
                ""
                if ideal is not None
                else "AI could not generate an ideal answer right now ("
                     + (ideal_error or "unavailable")
                     + "). Your answer was still saved."
            ),
        })
        return jsonify(result)

    filler_data = calculate_filler_words(text)
    star = detect_star(text)
    star_score = sum(star.values()) * 25

    # Run Gemini once. If it fails (quota, transient network error such as
    # WinError 10054, etc.) we degrade gracefully with real local metrics so the
    # answer is still saved and the user can continue the interview.
    ai_available = True
    ai_result = None
    ai_error_msg = None
    try:
        ai_result = analyze_with_gemini(text, interview_type, question, interview_role)
    except AIError as e:
        ai_available = False
        ai_error_msg = str(e)
        print(f"EVALUATE AI UNAVAILABLE: {ai_error_msg}")
    except Exception as e:
        # Never let an unexpected evaluation error break the interview flow.
        ai_available = False
        ai_error_msg = str(e)
        print("EVALUATE ERROR:", e)
        print(traceback.format_exc())
    if ai_result is not None and "error" in ai_result:
        ai_available = False
        ai_error_msg = ai_result.get("error")
        print(f"EVALUATE AI RESPONSE ERROR: {ai_error_msg}")

    if ai_available and ai_result is not None:
        validated_scores = validate_scores(ai_result.get("scores", {}))
        overall_score = calculate_overall_score(validated_scores)
    else:
        # Local, non-AI scoring: independent per-category heuristics so the
        # fallback never fabricates identical category scores.
        validated_scores = _local_fallback_scores(
            text, question, star_score, filler_data.get("rate") or 0
        )
        overall_score = calculate_overall_score(validated_scores)

    if ai_available and ai_result is not None:
        strengths = ai_result.get("strengths", [])
        areas_to_improve = ai_result.get("areas_to_improve", [])
        suggestions = ai_result.get("suggestions", [])
        feedback = ai_result.get("feedback", "") or ai_result.get("ai_feedback", "")
        detailed_analysis = ai_result.get("detailed_analysis", {})
        corrected_answer = ai_result.get("corrected_answer", "")
        improved_sample_answer = ai_result.get("improved_sample_answer", "")
        speaking_analysis = ai_result.get("speaking_analysis")
        interview_tip = ai_result.get("interview_tip", "")
        ai_error = ""
        feedback_note = ""
    else:
        # Honest degradation: no fake strengths/analysis. The real reason is
        # surfaced in `ai_error` / `feedback_note` and logged.
        strengths = []
        areas_to_improve = []
        suggestions = []
        feedback = ""
        detailed_analysis = {}
        corrected_answer = ""
        improved_sample_answer = ""
        speaking_analysis = None
        interview_tip = "Keep going — your answer has been saved."
        ai_error = ai_error_msg or "AI evaluation is currently unavailable."
        feedback_note = (
            "AI evaluation could not be completed right now (" + ai_error +
            "). Your answer was still saved."
        )

    result = {
        "success": True,
        "overall_score": overall_score,
        "scores": validated_scores,
        "clarity": validated_scores.get("clarity", 0),
        "communication": validated_scores.get("communication", 0),
        "confidence": validated_scores.get("confidence", 0),
        "grammar": validated_scores.get("grammar", 0),
        "relevance": validated_scores.get("relevance", 0),
        "answer_structure": validated_scores.get("structure", 0),
        "detailed_analysis": detailed_analysis,
        "strengths": strengths,
        "areas_to_improve": areas_to_improve,
        "improvements": areas_to_improve,
        "suggestions": suggestions,
        "corrected_answer": corrected_answer,
        "improved_sample_answer": improved_sample_answer,
        "feedback": feedback,
        "ai_feedback": feedback,
        "better_answer": improved_sample_answer,
        "interview_tip": interview_tip,
        "speaking_analysis": speaking_analysis,
        "ai_error": ai_error,
        "filler_data": filler_data,
        "star_analysis": star,
        "star_score": star_score,
        "ai_available": ai_available,
        "feedback_note": feedback_note
    }

    # NOTE: Individual answers are NOT persisted here. Each flow (practice,
    # mock-interview, resume-interview) aggregates scores and saves one
    # session-level record via /save-interview when the session completes.
    # This avoids double-counting and keeps stats accurate.

    return jsonify(result)


# ==========================================
# AI IMPROVEMENT PLAN
# ==========================================

def get_user_interview_history(user_id):
    conn = get_db()
    if not conn:
        return []
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """SELECT id, interview_type, overall_score, created_at,
                      communication, confidence, clarity, grammar, structure, relevance,
                      details, strengths, improvements, overall_feedback
               FROM interviews WHERE user_id = %s AND communication IS NOT NULL
               ORDER BY created_at DESC LIMIT 20""",
            (user_id,)
        )
        return cursor.fetchall()
    except mysql.connector.Error as e:
        print("Interview history error:", e)
        return []
    finally:
        cursor.close()
        conn.close()


def calculate_skill_stats(interviews):
    if not interviews:
        return {}
    cats = ["communication", "confidence", "clarity", "grammar", "structure", "relevance"]
    stats = {}
    for c in cats:
        vals = [r[c] for r in interviews if r.get(c) is not None]
        if vals:
            stats[c] = {
                "current": vals[0],
                "average": round(sum(vals) / len(vals)),
                "history": vals[:5],
                "min": min(vals),
                "max": max(vals),
                "count": len(vals),
                "trend": vals[0] - vals[1] if len(vals) > 1 else 0
            }
    return stats


def generate_ai_improvement_plan(interviews, skill_stats):
    if not client:
        return None
    if not interviews:
        return None

    latest = interviews[0]
    overall = latest["overall_score"]
    cats = ["communication", "confidence", "clarity", "grammar", "structure", "relevance"]
    cat_labels = {"communication": "Communication", "confidence": "Confidence",
                  "clarity": "Clarity", "grammar": "Grammar",
                  "structure": "Structure", "relevance": "Relevance"}

    scores_text = "\n".join([f"- {cat_labels[c]}: {latest.get(c, 'N/A')}/100" for c in cats])

    # Ground the plan in the candidate's REAL recorded weaknesses. Completed
    # mock interviews store aggregated `improvements` JSON plus per-question
    # `details` (answers, scores, AI feedback, improved answers). Feed those in
    # so the AI plan references actual weak topics instead of generic advice.
    weakness_lines = []
    feedback_lines = []
    improvements_json = _load_json_list(latest.get("improvements"))
    for item in improvements_json[:8]:
        t = str(item).strip()
        if t and t not in weakness_lines:
            weakness_lines.append(t)
    details = _load_json_list(latest.get("details"))
    for d in details[:10]:
        if isinstance(d, dict):
            fb = (d.get("feedback") or "").strip()
            if fb:
                feedback_lines.append(fb[:300])
            question_text = (d.get("question") or "").strip()
            score = d.get("score")
            if question_text and score is not None and float(score) < 70:
                weakness_lines.append(
                    "Low-scoring topic (%s/100): %s" % (score, question_text[:120])
                )
    recorded_weakness_context = ""
    if weakness_lines:
        recorded_weakness_context = (
            "\nReported weaknesses from your latest interview:\n- "
            + "\n- ".join(weakness_lines[:10])
            + "\n"
        )
    if feedback_lines:
        recorded_weakness_context += (
            "\nPer-question AI feedback notes from your latest interview:\n"
            + "\n".join("- " + line for line in feedback_lines[:5])
            + "\n"
        )

    history_text = ""
    if len(interviews) > 1:
        history_text = "\nPrevious interviews:\n"
        for i, inv in enumerate(interviews[:5]):
            history_text += f"  Interview {i+1}: Overall={inv['overall_score']}, "
            history_text += ", ".join([f"{cat_labels[c]}={inv.get(c, 'N/A')}" for c in cats])
            history_text += f" ({inv['interview_type']})\n"

    prompt = f"""You are an expert AI interview coach. Analyze the user's interview performance and create a personalized improvement plan.

Current Interview Scores (out of 100):
{scores_text}

Overall Score: {overall}/100
Interview Type: {latest['interview_type']}
Number of Interviews: {len(interviews)}
{recorded_weakness_context}{history_text}
Return ONLY a valid JSON object with these fields:
{{
    "strongest_skill": "<skill name>",
    "strongest_score": <number>,
    "weakest_skill": "<skill name>",
    "weakest_score": <number>,
    "ai_insight": "<2-3 sentences analyzing their weakest area with specific advice>",
    "recommendations": [
        {{"title": "<recommendation title>", "description": "<specific actionable advice>"}},
        {{"title": "<recommendation title>", "description": "<specific actionable advice>"}},
        {{"title": "<recommendation title>", "description": "<specific actionable advice>"}},
        {{"title": "<recommendation title>", "description": "<specific actionable advice>"}}
    ],
    "seven_day_plan": [
        {{"day": 1, "focus": "<skill name>", "activity": "<specific practice activity>", "goal": "<measurable goal>"}},
        {{"day": 2, "focus": "<skill name>", "activity": "<specific practice activity>", "goal": "<measurable goal>"}},
        {{"day": 3, "focus": "<skill name>", "activity": "<specific practice activity>", "goal": "<measurable goal>"}},
        {{"day": 4, "focus": "<skill name>", "activity": "<specific practice activity>", "goal": "<measurable goal>"}},
        {{"day": 5, "focus": "<skill name>", "activity": "<specific practice activity>", "goal": "<measurable goal>"}},
        {{"day": 6, "focus": "<skill name>", "activity": "<specific practice activity>", "goal": "<measurable goal>"}},
        {{"day": 7, "focus": "Full Interview", "activity": "<comprehensive practice activity>", "goal": "<measurable goal>"}}
    ],
    "summary": "<2-3 sentence personalized AI coach summary>"
}}

Scoring thresholds:
- 90-100: Excellent
- 80-89: Very Good
- 70-79: Good
- 60-69: Needs Improvement
- Below 60: Priority Area

Make recommendations specific to their actual weak areas. Focus 70% on weak areas, 30% on maintaining strengths.
Return ONLY the JSON object, no other text."""

    try:
        raw = gemini_generate_text(prompt, feature="improvement_plan")
        result = json.loads(raw)
        if not isinstance(result, dict):
            return None
        return result
    except AIError as e:
        print("Improvement plan AI error:", str(e))
        return None
    except Exception as e:
        print("Improvement plan AI error:", type(e).__name__, str(e)[:300])
        return None


@app.route("/improvement-plan", methods=["GET"])
@login_required
def improvement_plan():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    interviews = get_user_interview_history(user["id"])
    skill_stats = calculate_skill_stats(interviews) if interviews else {}
    return render_template("improvement_plan.html",
                           interviews=interviews,
                           skill_stats=skill_stats, plan=None, **user_context())


@app.route("/improvement-plan/generate", methods=["POST"])
@login_required
def generate_improvement_plan():
    user = get_current_user()
    if not user:
        return jsonify({"error": "Not authenticated"}), 401

    interviews = get_user_interview_history(user["id"])
    if not interviews:
        return jsonify({"error": "No interviews found"}), 400

    skill_stats = calculate_skill_stats(interviews)
    plan = generate_ai_improvement_plan(interviews, skill_stats)

    # If the AI plan is unavailable, still return the local skill_stats so the
    # frontend can compute strongest/weakest areas and show fallback feedback.
    if not plan:
        return jsonify({
            "plan": None,
            "skill_stats": skill_stats,
            "latest_overall": interviews[0]["overall_score"]
        })

    return jsonify({
        "plan": plan,
        "skill_stats": skill_stats,
        "latest_overall": interviews[0]["overall_score"]
    })


# ==========================================
# RESUME-BASED INTERVIEW
# ==========================================

UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), "uploads")
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_FILE_SIZE = 5 * 1024 * 1024

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text_from_pdf(filepath):
    text = ""
    try:
        with open(filepath, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
    except Exception as e:
        print("PDF extraction error:", e)
    return text.strip()


def extract_text_from_docx(filepath):
    text = ""
    try:
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
    except Exception as e:
        print("DOCX extraction error:", e)
    return text.strip()


def extract_resume_text(filepath, filetype):
    if filetype == "pdf":
        return extract_text_from_pdf(filepath)
    elif filetype == "docx":
        return extract_text_from_docx(filepath)
    elif filetype == "txt":
        try:
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read().strip()
        except Exception as e:
            print("TXT extraction error:", e)
            return ""
    return ""


def analyze_resume_with_ai(resume_text):
    if not client:
        return None
    prompt = f"""Analyze this resume and extract structured information. Return ONLY valid JSON.

Resume text:
{resume_text[:4000]}

Return a JSON object with:
{{
    "candidate_summary": "1-2 sentence professional summary",
    "skills": ["skill1", "skill2", ...],
    "programming_languages": ["lang1", "lang2", ...],
    "frameworks": ["framework1", ...],
    "databases": ["db1", ...],
    "projects": [
        {{"name": "Project Name", "description": "Brief description", "technologies": ["tech1", ...]}}
    ],
    "experience": [
        {{"title": "Job Title", "company": "Company", "duration": "Duration", "description": "Brief description"}}
    ],
    "education": [
        {{"degree": "Degree", "institution": "Institution", "year": "Year"}}
    ],
    "certifications": ["cert1", ...]
}}

Only include fields that have actual data from the resume. Return empty arrays for missing fields.
Return ONLY the JSON object."""
    try:
        raw = gemini_generate_text(prompt, feature="resume_analysis")
        result = json.loads(raw)
        if not isinstance(result, dict):
            raise AIError("AI did not return a valid analysis object.", category="server")
        return result
    except (AIError,) as e:
        raise e
    except Exception as e:
        print("Resume analysis error:", type(e).__name__, str(e)[:300])
        raise AIError(str(e), category="server") from e


RESUME_DIFF_LABELS = {
    "easy": "Beginner",
    "beginner": "Beginner",
    "medium": "Intermediate",
    "intermediate": "Intermediate",
    "hard": "Advanced",
    "advanced": "Advanced",
}

RESUME_TYPE_LABELS = {
    "hr": "an HR / cultural-fit interview",
    "technical": "a technical interview",
    "behavioral": "a behavioral (STAR method) interview",
    "communication": "a communication-skills interview",
    "general": "a general interview",
    "mixed": "a mixed interview",
}

RESUME_EXP_LABELS = {
    "fresher": "Fresher (0 years)",
    "junior": "Junior (1-3 years)",
    "mid": "Mid (3-6 years)",
    "senior": "Senior (6+ years)",
}


def guess_question_category(text):
    """Best-effort category label for a generated question."""
    low = (text or "").lower()
    technical_kw = (
        "python", "java", "javascript", "typescript", "c++", "code", "coding",
        "api", "rest", "database", "sql", "nosql", "react", "angular", "vue",
        "django", "flask", "spring", "node", "express", "framework", "library",
        "function", "algorithm", "algorithm", "data structure", "architecture",
        "oop", "object-oriented", "test", "testing", "debug", "deployment",
        "deploy", "git", "version control", "linux", "aws", "docker", "kubernetes",
        "technology", "programming", "language", "machine learning", "data analysis",
        "pandas", "numpy", "html", "css", "frontend", "backend", "server",
    )
    project_kw = ("project", "built", "developed", "implemented", "design", "system", "module")
    if any(k in low for k in project_kw):
        return "Project"
    if any(k in low for k in technical_kw):
        return "Technical"
    if any(k in low for k in (
            "experience", "team", "collaborat", "conflict", "lead", "leadership",
            "behavior", "tell me about a time", "you handled", "worked with")):
        return "Behavioral"
    if any(k in low for k in ("certif", "education", "degree", "course", "university", "learn", "study")):
        return "Education"
    return "General"


def normalize_question_objects(questions, difficulty):
    """Normalize a list of raw questions (strings or dicts) into the
    standardized response shape:
        {"id": 1, "question": "...", "category": "...", "difficulty": "..."}
    Deduplicates and drops empty placeholders so the count is accurate."""
    diff_label = RESUME_DIFF_LABELS.get((difficulty or "medium").strip().lower(), "Medium")
    result = []
    seen = set()
    idx = 0
    for q in questions:
        if isinstance(q, dict):
            text = str(q.get("question") or q.get("text") or q.get("content") or "").strip()
            category = str(q.get("category") or guess_question_category(text))
            diff = str(q.get("difficulty") or diff_label)
        else:
            text = str(q).strip()
            category = guess_question_category(text)
            diff = diff_label
        if not text or len(text) < 5:
            continue
        key = text.lower()
        if key in seen:
            continue
        seen.add(key)
        idx += 1
        result.append({"id": idx, "question": text, "category": category, "difficulty": diff})
    return result


def _coerce_questions(obj, depth=0):
    """Walk a parsed JSON value and return the first list of question-like
    items (a JSON array, or a nested "questions" / array inside a wrapper
    object). Falls back to the value itself so callers can still validate."""
    if depth > 8 or not isinstance(obj, (dict, list)):
        return obj
    if isinstance(obj, list):
        return obj
    candidates = []
    if isinstance(obj.get("questions"), list):
        candidates.append(obj["questions"])
    for value in obj.values():
        if isinstance(value, (dict, list)):
            candidates.append(value)
    for candidate in candidates:
        result = _coerce_questions(candidate, depth + 1)
        if isinstance(result, list):
            return result
    return obj


def _extract_json_blocks(raw):
    """Robustly parse AI text into a JSON value.

    Tolerates markdown code fences, prose wrapped around the JSON, and nested
    wrappers such as {"questions": [...]} or {"data": {"questions": [...]}}.
    Raises AIError when nothing usable can be extracted so the route can
    surface the real reason instead of silently returning an empty list.
    """
    text = str(raw or "").strip()
    if not text:
        raise AIError("AI returned an empty response.", category="server")

    attempts = [text]

    # Strip a markdown code fence, if present.
    stripped = re.sub(r"^```[a-zA-Z]*\s*", "", text)
    stripped = re.sub(r"\s*```$", "", stripped).strip()
    if stripped != text:
        attempts.append(stripped)

    # Slice the first '[' .. last ']' span (array wrapped in prose).
    arr_start = text.find("[")
    arr_end = text.rfind("]")
    if arr_start != -1 and arr_end > arr_start:
        attempts.append(text[arr_start:arr_end + 1])

    # Slice the first '{' .. last '}' span (object wrapped in prose).
    obj_start = text.find("{")
    obj_end = text.rfind("}")
    if obj_start != -1 and obj_end > obj_start:
        attempts.append(text[obj_start:obj_end + 1])

    for candidate in attempts:
        try:
            parsed = json.loads(candidate)
        except (ValueError, TypeError):
            continue
        return _coerce_questions(parsed)

    raise AIError(
        "AI returned a response that could not be parsed as questions.",
        category="server",
    )


def generate_resume_questions_gemini(
    resume_data,
    role,
    experience,
    interview_type,
    difficulty,
    count,
    resume_text="",
):
    """Generate exactly `count` personalized questions using the configured
    Gemini client (GEMINI_API_KEY). Questions are grounded in the candidate's
    actual resume skills, languages, frameworks, projects, education,
    certifications, experience and the requested target role/difficulty.

    Raises AIError on failure. Never returns an empty list.
    """
    count = max(1, min(int(count), 20))

    skills = ", ".join(str(s) for s in resume_data.get("skills", [])[:12]) or "Not listed"
    langs = ", ".join(str(s) for s in resume_data.get("programming_languages", [])[:8]) or "Not listed"
    frameworks = ", ".join(str(s) for s in resume_data.get("frameworks", [])[:6]) or "Not listed"
    databases = ", ".join(str(s) for s in resume_data.get("databases", [])[:5]) or "Not listed"
    certifications = ", ".join(str(s) for s in resume_data.get("certifications", [])[:6]) or "Not listed"
    projects = "\n".join([
        "- %s: %s (Technologies: %s)"
        % (
            p.get("name", "Unknown"),
            str(p.get("description", ""))[:120],
            ", ".join(str(t) for t in p.get("technologies", [])[:6]),
        )
        for p in resume_data.get("projects", [])[:5]
    ]) or "None listed"
    experience_text = "\n".join([
        "- %s at %s (%s)"
        % (e.get("title", ""), e.get("company", ""), e.get("duration", ""))
        for e in resume_data.get("experience", [])[:4]
    ]) or "None listed - Fresher"
    education = ", ".join([
        "%s from %s%s"
        % (e.get("degree", ""), e.get("institution", ""),
           (" (%s)" % e.get("year", "")) if e.get("year") else "")
        for e in resume_data.get("education", [])[:3]
    ]) or "Not listed"

    candidate_summary = str(resume_data.get("candidate_summary") or "").strip()
    raw_resume_excerpt = str(resume_text or "").strip()[:1500]

    type_label = RESUME_TYPE_LABELS.get((interview_type or "mixed").strip().lower(), RESUME_TYPE_LABELS["mixed"])
    diff_label = RESUME_DIFF_LABELS.get((difficulty or "medium").strip().lower(), "Intermediate")
    exp_label = RESUME_EXP_LABELS.get((experience or "fresher").strip().lower(),
                                      str(experience or "Fresher").replace("_", " "))

    questions = []
    seen = set()
    max_calls = 4

    def generate_round(existing):
        want = count - len(existing)
        if want <= 0:
            return []
        existing_block = ""
        if existing:
            existing_block = (
                "Questions already created (do NOT repeat or rephrase ANY of them):\n"
                + "\n".join("- " + q for q in existing)
                + "\n\n"
            )
        prompt = f"""You are an expert technical interviewer and hiring coach. Create exactly {want} personalized interview questions for the candidate below.

Resume Summary:
{candidate_summary or "Not provided"}

Raw Resume Excerpt:
{raw_resume_excerpt or "Not provided"}

Structured Resume Information:
- Skills: {skills}
- Programming Languages: {langs}
- Frameworks: {frameworks}
- Databases: {databases}
- Certifications: {certifications}
- Projects:
{projects}
- Work Experience:
{experience_text}
- Education: {education}

Interview Configuration:
- Target Role: {role}
- Experience Level: {exp_label}
- Interview Type: {type_label}
- Difficulty: {diff_label}
- Number of Questions Needed Right Now: {want}

Requirements:
- EVERY question MUST reference the candidate's ACTUAL resume content: named skills, programming languages, frameworks, databases, projects, education, certifications or experience. Do NOT write generic questions that ignore the resume.
- For resume projects, ask about architecture, design decisions, technologies used, challenges, testing, or improvements.
- For skills/languages/frameworks, ask about practical application, depth, and trade-offs in real work.
- Tailor difficulty and depth to a {diff_label}-level {role} candidate at the {exp_label} experience level.
- Mix the requested interview type ({type_label}). For "mixed", combine technical, project, behavioral and follow-up questions.
- Questions must be unique, specific, realistic and thought-provoking. Progress from slightly easier to harder.
- Do NOT include hints, answers, explanations, bullet labels, or conversation.
{existing_block}- Return ONLY a valid JSON array of exactly {want} question objects, each with "question", "category" (e.g. "Technical", "Project", "Behavioral", "Education", "General") and "difficulty" fields, e.g.
[{{"question": "Describe the Flask project you built and how you handled user authentication.", "category": "Project", "difficulty": "Intermediate"}}]"""
        raw = gemini_generate_text(prompt, feature="resume_question_generation")
        return _extract_json_blocks(raw)

    for _ in range(max_calls):
        parsed = generate_round(questions)
        items = []
        if isinstance(parsed, dict):
            if isinstance(parsed.get("questions"), list):
                items = parsed["questions"]
            else:
                for value in parsed.values():
                    if isinstance(value, list):
                        items = value
                        break
        elif isinstance(parsed, list):
            items = parsed
        for item in items:
            if isinstance(item, str):
                text = str(item).strip()
            elif isinstance(item, dict):
                text = str(item.get("question") or item.get("text") or item.get("content") or "").strip()
            else:
                continue
            if text and len(text) > 5 and text.lower() not in seen:
                seen.add(text.lower())
                questions.append(text)
                if len(questions) >= count:
                    break
        if len(questions) >= count:
            break

    if not questions:
        raise AIError("AI returned no usable questions.", category="server")
    return questions[:count]


def generate_resume_questions(resume_data, role, experience, interview_type, difficulty, count, resume_text=""):
    """Generate resume-personalized questions.

    Uses the configured Gemini client first (the same client the rest of the
    app uses every day), then falls back to the OpenAI-based generator for
    compatibility. NEVER returns None silently: on total failure it raises
    AIError with the real reason so the API can surface it to the user.
    """
    error_messages = []

    if client:
        try:
            questions = generate_resume_questions_gemini(
                resume_data, role, experience, interview_type, difficulty, count, resume_text
            )
            if questions:
                print("Resume question generation succeeded via Gemini; count=%d" % len(questions))
                return questions
        except AIError as e:
            error_messages.append("Gemini: " + str(e))
            print("Resume question generation (Gemini) failed:", str(e))
        except Exception as e:
            error_messages.append("Gemini: " + str(e))
            print("Resume question generation (Gemini) error:", type(e).__name__, str(e)[:300])

    try:
        questions = generate_resume_questions_ai(
            resume_data, role, experience, interview_type, difficulty, count
        )
        if questions:
            print("Resume question generation succeeded via OpenAI; count=%d" % len(questions))
            return questions
    except Exception as e:
        error_messages.append("OpenAI: " + str(e))
        print("Resume question generation (OpenAI) failed:", type(e).__name__, str(e)[:300])

    if error_messages:
        raise AIError(
            "Question generation failed: " + " | ".join(error_messages),
            category="config" if not client else "server",
        )
    raise AIError("No AI provider is configured (check GEMINI_API_KEY / OPENAI_API_KEY in .env).", category="config")


@app.route("/resume")
@login_required
def resume_upload():
    role_display = {}
    for k, v in ROLE_DESCRIPTIONS.items():
        role_display[k] = v.split(" with ")[0].split(" proficient")[0].split(" specializing")[0]
    return render_template("resume_upload.html", roles=role_display, **user_context())


@app.route("/api/resume/analyze", methods=["POST"])
@login_required
def api_analyze_resume():
    if "resume_file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400

    file = request.files["resume_file"]
    if file.filename == "":
        return jsonify({"error": "No file selected."}), 400

    if not allowed_file(file.filename):
        return jsonify({"error": "Please upload a valid PDF, DOCX or TXT file."}), 400

    filename = secure_filename(file.filename)
    ext = filename.rsplit(".", 1)[1].lower()
    unique_name = f"{uuid.uuid4().hex[:8]}_{filename}"
    filepath = os.path.join(UPLOAD_FOLDER, unique_name)
    file.save(filepath)

    if os.path.getsize(filepath) > MAX_FILE_SIZE:
        os.remove(filepath)
        return jsonify({"error": "File too large. Maximum size is 5 MB."}), 400

    resume_text = extract_resume_text(filepath, ext)
    os.remove(filepath)

    if not resume_text or len(resume_text) < 20:
        return jsonify({"error": "Unable to extract text from this resume. Please upload a text-based PDF or DOCX."}), 400

    # Session-level cache: if this exact resume text was already analyzed in this
    # session, reuse the result instead of calling Gemini again.
    resume_hash = hashlib.sha256(resume_text.encode("utf-8")).hexdigest()[:16]
    cache = session.get("resume_analysis_cache", {})
    if resume_hash in cache:
        entry = cache[resume_hash]
        if isinstance(entry, dict) and "analysis" in entry:
            analysis = entry["analysis"]
            cached_text = entry.get("text") or resume_text
        else:
            analysis = entry
            cached_text = resume_text
        _log_ai_call("resume_analysis_cached", True, status=200, detail="reused cached analysis")
        return jsonify({"success": True, "analysis": analysis,
                        "resume_text": cached_text,
                        "text_length": len(cached_text), "cached": True})

    try:
        analysis = analyze_resume_with_ai(resume_text)
    except AIError as e:
        return jsonify({"success": False, "error": str(e), "error_type": e.category, "questions": []}), 502

    # Keep only the most recent few analyses to avoid unbounded session growth.
    cache[resume_hash] = {"analysis": analysis, "text": resume_text}
    if len(cache) > 3:
        # drop the oldest entry
        cache.pop(next(iter(cache)))
    session["resume_analysis_cache"] = cache

    return jsonify({"success": True, "analysis": analysis,
                    "resume_text": resume_text, "text_length": len(resume_text), "cached": False})


@app.route("/api/resume/generate-questions", methods=["POST"])
@login_required
def api_generate_resume_questions():
    data = request.get_json(silent=True) or {}
    resume_data = data.get("resume_data") or {}
    resume_text = data.get("resume_text") or ""
    role = data.get("role", "software_developer")
    experience = data.get("experience", "fresher")
    interview_type = data.get("interview_type", "mixed")
    difficulty = data.get("difficulty", "medium")

    # Backend-safe count validation: default 10, min 1, max 20.
    try:
        count = max(1, min(int(data.get("count", 10)), 20))
    except (TypeError, ValueError):
        count = 10

    print("RESUME Q GEN REQUEST: role=%s experience=%s type=%s difficulty=%s count=%s "
          "resume_data_keys=%s resume_text_length=%s"
          % (role, experience, interview_type, difficulty, count,
             list(resume_data.keys()) if isinstance(resume_data, dict) else "N/A",
             len(resume_text or "")))

    if not isinstance(resume_data, dict) or not resume_data:
        has_content = False
    else:
        # Confirm the resume actually contains usable content to personalize
        # against (skills, languages, frameworks, projects, experience,
        # certifications, or raw resume text).
        has_content = (
            resume_data.get("skills")
            or resume_data.get("programming_languages")
            or resume_data.get("frameworks")
            or resume_data.get("databases")
            or resume_data.get("projects")
            or resume_data.get("experience")
            or resume_data.get("certifications")
            or bool((resume_text or "").strip())
        )

    # Validate BEFORE calling the AI/generation logic. Never generate questions
    # against an empty resume.
    if not has_content:
        return jsonify({
            "success": False,
            "error": "No resume information found. Please upload or complete your resume first.",
            "error_type": "validation",
            "questions": [],
        }), 400

    role_desc = ROLE_DESCRIPTIONS.get(role, role)
    try:
        questions = generate_resume_questions(
            resume_data, role_desc, experience, interview_type, difficulty, count, resume_text
        )
    except AIError as e:
        print("RESUME Q GEN ERROR:", str(e))
        return jsonify({
            "success": False,
            "error": str(e),
            "error_type": e.category,
            "questions": [],
        }), 502

    if not questions:
        return jsonify({
            "success": False,
            "error": "Question generation returned no questions. Please try again.",
            "error_type": "server",
            "questions": [],
        }), 502

    normalized = normalize_question_objects(questions, difficulty)
    print("RESUME Q GEN RESPONSE: success=yes count=%s" % len(normalized))
    return jsonify({
        "success": True,
        "questions": normalized,
        "count": len(normalized),
        "role": role,
        "experience": experience,
        "difficulty": difficulty,
        "interview_type": interview_type,
    })


@app.route("/api/resume/process", methods=["POST"])
@login_required
def api_process_resume():
    """Single-call flow: upload resume -> extract text -> analyze -> generate questions."""
    if "resume_file" not in request.files:
        return jsonify({"error": "No file uploaded. Please choose a resume."}), 400

    file = request.files["resume_file"]
    if file.filename == "":
        return jsonify({"error": "No file selected. Please choose a resume."}), 400

    if not allowed_file(file.filename):
        return jsonify({"error": "Unsupported file type. Please upload a PDF, DOCX or TXT resume."}), 400

    filename = secure_filename(file.filename)
    ext = filename.rsplit(".", 1)[1].lower()
    unique_name = f"{uuid.uuid4().hex[:8]}_{filename}"
    filepath = os.path.join(UPLOAD_FOLDER, unique_name)
    file.save(filepath)

    if os.path.getsize(filepath) > MAX_FILE_SIZE:
        os.remove(filepath)
        return jsonify({"error": "File too large. Maximum size is 5 MB."}), 400

    resume_text = extract_resume_text(filepath, ext)
    os.remove(filepath)

    if not resume_text or len(resume_text) < 20:
        return jsonify({"error": "Could not read text from this resume. Please upload a text-based PDF or DOCX."}), 400

    role = request.form.get("role", "software_developer")
    experience = request.form.get("experience", "fresher")
    interview_type = request.form.get("interview_type", "mixed")
    difficulty = request.form.get("difficulty", "medium")
    try:
        count = max(1, min(int(request.form.get("count", 10)), 20))
    except (TypeError, ValueError):
        count = 10

    try:
        analysis = analyze_resume_with_ai(resume_text)
        if not analysis:
            return jsonify({"success": False, "error": "Resume analysis could not be completed. Please try again.",
                            "error_type": "server", "questions": []}), 502
        role_desc = ROLE_DESCRIPTIONS.get(role, role)
        questions = generate_resume_questions(analysis, role_desc, experience, interview_type, difficulty, count, resume_text)
    except AIError as e:
        print("Resume process error:", str(e))
        return jsonify({"success": False, "error": str(e), "error_type": e.category, "questions": []}), 502

    if not questions or len(questions) == 0:
        return jsonify({"success": False,
                        "error": "Question generation returned no questions. Please try again.",
                        "error_type": "server", "questions": []}), 502

    normalized = normalize_question_objects(questions, difficulty)

    # Persist the generated interview state server-side (per user) so the
    # interview survives page refresh / re-entry / new tab on the same account -
    # without re-uploading the resume or re-calling Gemini.
    user = get_current_user()
    if user:
        save_resume_interview_state(user["id"], normalized, analysis, role, experience, len(normalized))

    return jsonify({
        "success": True,
        "analysis": analysis,
        "questions": normalized,
        "count": len(normalized),
        "text_length": len(resume_text),
        "role": role,
        "experience": experience,
        "difficulty": difficulty,
        "interview_type": interview_type,
    })


@app.route("/resume-interview")
@login_required
def resume_interview():
    role_display = {}
    for k, v in ROLE_DESCRIPTIONS.items():
        role_display[k] = v.split(" with ")[0].split(" proficient")[0].split(" specializing")[0]
    return render_template("resume_interview.html", roles=role_display, **user_context())


@app.route("/api/resume/interview-state", methods=["GET"])
@login_required
def resume_interview_state_get():
    """Return the server-persisted interview state (questions + analysis) so the
    page can restore an in-progress interview without re-uploading or calling
    Gemini again. Returns empty when there is no active state."""
    user = get_current_user()
    if not user:
        return jsonify({"success": True, "has_state": False})
    data = get_resume_interview_state(user["id"])
    if not data or not data.get("questions"):
        return jsonify({"success": True, "has_state": False})
    return jsonify({
        "success": True,
        "has_state": True,
        "questions": data["questions"],
        "analysis": data.get("analysis"),
        "count": len(data["questions"])
    })


@app.route("/api/resume/interview-state", methods=["DELETE"])
@login_required
def resume_interview_state_delete():
    """Clear the server-persisted interview state (used by Start New Interview)."""
    user = get_current_user()
    if user:
        delete_resume_interview_state(user["id"])
    return jsonify({"success": True})


@app.route("/save-interview", methods=["POST"])
@login_required
def save_interview():
    user = get_current_user()
    if not user:
        return jsonify({"error": "Not authenticated"}), 401

    data = request.get_json(silent=True) or {}
    interview_type = data.get("interview_type", "mixed")
    overall_score = int(data.get("overall_score", 70))
    questions_count = int(data.get("questions_count", 0))
    communication = int(data.get("communication", 70))
    confidence = int(data.get("confidence", 70))
    clarity = int(data.get("clarity", 70))
    grammar = int(data.get("grammar", 70))
    structure = int(data.get("structure", 70))
    relevance = int(data.get("relevance", 70))

    print(f"SAVE INTERVIEW: user_id={user['id']} type={interview_type} "
          f"overall={overall_score} questions={questions_count}")

    # Optional report enrichment fields persisted with the interview so the
    # Reports page can rebuild the full detail view from stored data only
    # (no Gemini call is ever made when opening / refreshing Reports).
    details = data.get("details") or []
    strengths = data.get("strengths") or []
    improvements = data.get("improvements") or []
    role = (data.get("role") or "").strip()[:100] or None
    experience = (data.get("experience") or "").strip()[:50] or None
    difficulty = (data.get("difficulty") or "").strip()[:20] or None
    overall_feedback = (data.get("overall_feedback") or "").strip() or None

    def _as_json_list(v):
        return json.dumps(v, default=str) if isinstance(v, list) and v else None

    # Classify the session (practice / mock / resume) so admin analytics can
    # count session types reliably. The practice, mock and resume interview
    # pages all reuse `interview_type` values like 'hr'/'technical', so we
    # disambiguate using the distinctive overall_feedback text they send.
    def _infer_session_type(it, fb):
        fb_low = (fb or "").lower()
        if it in ("resume", "resume-based") or "resume" in it:
            return "resume"
        if "mock interview" in fb_low:
            return "mock"
        if "practice session" in fb_low:
            return "practice"
        return "practice"

    session_type = data.get("session_type") or _infer_session_type(interview_type, overall_feedback)
    if session_type not in ("practice", "mock", "resume"):
        session_type = "practice"

    conn = get_db()
    if not conn:
        return jsonify({"success": False, "error": "Database unavailable"}), 500
    try:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO interviews
            (user_id, interview_type, overall_score, questions_count,
             communication, confidence, clarity, grammar, structure, relevance,
             details, strengths, improvements, role, experience, difficulty, overall_feedback,
             session_type)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (
            user["id"], interview_type, overall_score, questions_count,
            communication, confidence, clarity, grammar, structure, relevance,
            _as_json_list(details), _as_json_list(strengths), _as_json_list(improvements),
            role, experience, difficulty, overall_feedback,
            session_type,
        ))
        conn.commit()
        print(f"SAVE INTERVIEW: db commit successful (user_id={user['id']}, "
              f"type={interview_type}, overall={overall_score}, questions={questions_count})")
    except Exception as e:
        print("Save interview error:", e)
        print(traceback.format_exc())
        return jsonify({"success": False, "error": "Unable to save interview."}), 500
    finally:
        cursor.close()
        conn.close()

    return jsonify({"success": True, "interview_type": interview_type,
                    "overall_score": overall_score, "questions_count": questions_count})


# ==========================================
# ADMIN ROUTES
# ==========================================

@app.route("/admin/dashboard")
@admin_required
def admin_dashboard():
    stats = get_admin_summary_stats()
    recent_users = get_recent_registrations(limit=5)
    recent_interviews = get_recent_interviews_all(limit=5)
    return render_template("admin_dashboard.html", stats=stats,
                           recent_users=recent_users,
                           recent_interviews=recent_interviews,
                           active_page="dashboard", **user_context())


@app.route("/admin-dashboard")
@admin_required
def admin_dashboard_shortcut():
    return redirect(url_for("admin_dashboard"))


@app.route("/admin/users")
@admin_required
def admin_users():
    page = request.args.get("page", 1, type=int)
    per_page = 20
    users = get_admin_all_users()
    total = len(users)
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    start = (page - 1) * per_page
    page_users = users[start:start + per_page]
    return render_template("admin_users.html", users=page_users, total=total,
                           page=page, total_pages=total_pages, per_page=per_page,
                           active_page="users", **user_context())


@app.route("/admin/user/<int:user_id>")
@admin_required
def admin_user_detail(user_id):
    data = get_admin_user_detail(user_id)
    target = data.get("user")
    if not target:
        return render_template("admin_user_detail.html", target_user=None, error="User not found.",
                               active_page="users", **user_context()), 404
    return render_template("admin_user_detail.html", target_user=target,
                           stats=data.get("stats"),
                           interviews=data.get("interviews"),
                           strongest=data.get("strongest"),
                           weakest=data.get("weakest"),
                           active_page="users", **user_context())


@app.route("/admin/user/<int:user_id>/toggle-status", methods=["POST"])
@admin_required
def admin_toggle_user_status(user_id):
    target = get_user_by_id(user_id)
    if not target:
        return jsonify({"error": "User not found."}), 404
    if target["id"] == session.get("user_id"):
        return jsonify({"error": "You cannot deactivate your own account."}), 400

    conn = get_db()
    if not conn:
        return jsonify({"error": "Database unavailable."}), 500
    cursor = None
    try:
        cursor = conn.cursor()
        new_status = 0 if target["is_active"] else 1
        cursor.execute("UPDATE users SET is_active = %s WHERE id = %s", (new_status, user_id))
        conn.commit()
        return jsonify({"success": True, "is_active": bool(new_status)})
    except mysql.connector.Error as e:
        print("Toggle user status error:", e)
        return jsonify({"error": "Unable to update user status."}), 500
    finally:
        if cursor is not None:
            try:
                cursor.close()
            except Exception:
                pass
        conn.close()


@app.route("/admin/user/<int:user_id>/role", methods=["POST"])
@admin_required
def admin_change_role(user_id):
    target = get_user_by_id(user_id)
    if not target:
        return jsonify({"error": "User not found."}), 404
    if target["id"] == session.get("user_id"):
        return jsonify({"error": "You cannot change your own role."}), 400

    new_role = (request.form.get("role") if request.form else None) or \
               ((request.json or {}).get("role") if request.is_json else None)
    if new_role not in ("user", "admin"):
        return jsonify({"error": "Invalid role."}), 400

    conn = get_db()
    if not conn:
        return jsonify({"error": "Database unavailable."}), 500
    cursor = None
    try:
        # Protect against demoting the last admin in the system.
        if target["role"] == "admin" and new_role != "admin":
            check = conn.cursor()
            check.execute("SELECT COUNT(*) FROM users WHERE role = 'admin' AND is_active = 1")
            admin_count = check.fetchone()[0]
            check.close()
            if admin_count <= 1:
                return jsonify({"error": "Cannot demote the last admin. Promote another admin first."}), 400
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET role = %s WHERE id = %s", (new_role, user_id))
        conn.commit()
        return jsonify({"success": True, "role": new_role})
    except mysql.connector.Error as e:
        print("Change role error:", e)
        return jsonify({"error": "Unable to update role."}), 500
    finally:
        if cursor is not None:
            try:
                cursor.close()
            except Exception:
                pass
        conn.close()


@app.route("/admin/analytics")
@admin_required
def admin_analytics():
    data = get_admin_analytics()
    summary = get_admin_summary_stats()
    if not data:
        data = {
            "cat_avgs": {key: {"avg": None, "count": 0} for key, _label in ADMIN_CATEGORIES},
            "common_weak": (None, 0),
            "common_strong": (None, 0),
            "activity": [],
            "registrations": [],
        }
    return render_template("admin_analytics.html", data=data, summary=summary,
                           categories=ADMIN_CATEGORIES, active_page="analytics",
                           **user_context())


@app.route("/admin/reports")
@admin_required
def admin_reports():
    """System reports: every interview across all users, searchable."""
    q = (request.args.get("q") or "").strip()
    page = request.args.get("page", 1, type=int)

    conn = get_db()
    if not conn:
        return render_template("admin_reports.html", rows=[], total=0, page=1,
                               total_pages=1, q=q, active_page="reports",
                               **user_context()), 500
    cursor = None
    rows = []
    try:
        cursor = conn.cursor(dictionary=True)
        where = ""
        params = []
        if q:
            where = ("WHERE i.interview_type LIKE %s OR u.full_name LIKE %s "
                     "OR u.email LIKE %s OR i.role LIKE %s")
            like = "%" + q + "%"
            params = [like, like, like, like]
        cursor.execute(
            """SELECT i.id, i.interview_type, i.overall_score, i.role, i.created_at,
                      u.id AS user_id, u.full_name, u.email
               FROM interviews i JOIN users u ON u.id = i.user_id
               """ + where + " ORDER BY i.created_at DESC",
            params
        )
        rows = cursor.fetchall()
    except mysql.connector.Error as e:
        print("Admin reports error:", e)
    finally:
        if cursor is not None:
            try:
                cursor.close()
            except Exception:
                pass
        conn.close()

    total = len(rows)
    per_page = 25
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    page_rows = rows[(page - 1) * per_page:page * per_page]
    return render_template("admin_reports.html", rows=page_rows, total=total,
                           page=page, total_pages=total_pages, q=q,
                           active_page="reports", **user_context())


# ==========================================
# ADMIN CREATION (Flask CLI)
# ==========================================

@app.cli.command("create-admin")
def create_admin_cli():
    """Create the first admin account from the CLI.

    Usage:
        flask create-admin
    """
    import click
    full_name = click.prompt("Full name", default="Admin")
    email = click.prompt("Email")
    password = click.prompt("Password", hide_input=True, confirmation_prompt=True)

    if len(password) < 8:
        click.echo("ERROR: Password must be at least 8 characters.")
        return

    conn = get_db()
    if not conn:
        click.echo("ERROR: Could not connect to the database.")
        return
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM users WHERE email = %s", (email.lower().strip(),))
        existing = cursor.fetchone()
        if existing:
            cursor.execute(
                "UPDATE users SET role = 'admin', is_active = 1 WHERE id = %s", (existing[0],)
            )
            click.echo("Existing user promoted to admin: %s" % email)
        else:
            password_hash = generate_password_hash(password)
            cursor.execute(
                "INSERT INTO users (full_name, email, password_hash, role, is_active) "
                "VALUES (%s, %s, %s, 'admin', 1)",
                (full_name.strip(), email.lower().strip(), password_hash)
            )
            click.echo("Admin account created: %s" % email)
        conn.commit()
    except mysql.connector.Error as e:
        print("Create admin error:", e)
        click.echo("ERROR: Could not create admin account.")
    finally:
        cursor.close()
        conn.close()


# ==========================================
# RUN
# ==========================================

if __name__ == "__main__":
    init_db()
    app.run(debug=True)
